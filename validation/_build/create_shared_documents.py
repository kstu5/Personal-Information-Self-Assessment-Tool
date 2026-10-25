"""Create synthetic, shareable document examples and an independent answer key."""
from pathlib import Path
import json,zipfile
from lxml import etree as E
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'examples/documents';OUT.mkdir(exist_ok=True)
NAMES=['김민수','이지은','박서윤','김하늘','최지우','정수빈','한지민','이도윤','윤서아','오지훈']
ADDRESSES=['서울특별시 마포구 예시로 12, 302호','경기도 수원시 영통구 가상로 24, 101동 502호','부산광역시 해운대구 예시길 18, 202호','대전광역시 서구 가상로 30, 401호','광주광역시 북구 예시로 42, 603호']
PEOPLE=[{'number':i+1,'name':n,'phone':f'010-0000-{i+1:04d}','email':f'winner{i+1:02d}@example.invalid','address':ADDRESSES[i%5],'prize':'도서상품권' if i<3 else '기념품 세트' if i<6 else '커피 교환권'} for i,n in enumerate(NAMES)]
def font(run,size=None,bold=None):
    run.font.name='맑은 고딕';run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
    if size is not None:run.font.size=Pt(size)
    if bold is not None:run.bold=bold
    run.font.color.rgb=RGBColor(0,0,0)
def base(title):
    d=Document();s=d.sections[0];s.page_width=Mm(210);s.page_height=Mm(297);s.top_margin=Mm(16);s.bottom_margin=Mm(16);s.left_margin=Mm(18);s.right_margin=Mm(18)
    d.core_properties.author='';d.core_properties.last_modified_by='';d.core_properties.title=title;d.core_properties.subject=''
    for name in ('Normal','Title','Heading 1','Heading 2'):
        st=d.styles[name];st.font.name='맑은 고딕';st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
        for e in list(st.element.iter()):
            if E.QName(e).localname=='pBdr':e.getparent().remove(e)
    normal=d.styles['Normal'];normal.font.size=Pt(10);normal.paragraph_format.space_after=Pt(4);normal.paragraph_format.line_spacing=1.1
    d.styles['Title'].font.size=Pt(21);d.styles['Title'].font.bold=True;d.styles['Title'].paragraph_format.space_after=Pt(4)
    h=d.styles['Heading 1'];h.font.size=Pt(11.5);h.font.bold=True;h.paragraph_format.space_before=Pt(8);h.paragraph_format.space_after=Pt(4)
    d.add_paragraph(title,'Title');p=d.add_paragraph('파일 검사 연습용 가상 문서  모든 인적사항과 기관명은 예시입니다.')
    for r in p.runs:font(r,8.5)
    return d
def para(d,text,bold=False):
    p=d.add_paragraph();font(p.add_run(text),bold=bold);return p
def heading(d,text):return d.add_paragraph(text,'Heading 1')
def table(d,headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,w in enumerate(widths):t.columns[i].width=Mm(w)
    for i,text in enumerate(headers):t.rows[0].cells[i].text=str(text)
    for row in rows:
        cells=t.add_row().cells
        for i,text in enumerate(row):cells[i].text=str(text)
    pr=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    for rowno,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        for i,c in enumerate(row.cells):
            c.width=Mm(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for edge,val in [('top','65'),('bottom','65'),('left','90'),('right','90')]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');m.append(e)
            tcpr.append(m)
            if rowno==0 or rowno%2==0:
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E8EEF4' if rowno==0 else 'F7F9FB');tcpr.append(fill)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
                for r in p.runs:font(r,9.5,rowno==0)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    return t
def save(d,name):
    for p in d.paragraphs:
        for r in p.runs:font(r)
    d.save(OUT/name)

def create_documents():
    d=base('개인정보 검사 예시 이력서');para(d,'지원 분야  데이터 운영 및 업무 지원',True)
    heading(d,'기본 정보')
    table(d,['항목','내용','항목','내용'],[
      ['성명','김민수','생년월일','1998-03-14'],['나이','28세','휴대전화','010-0000-1234'],
      ['이메일','minsu.kim@example.invalid','주소',ADDRESSES[0]]],[18,67,19,70])
    heading(d,'학력')
    table(d,['기간','학교','전공'],[['2017.03 ~ 2021.02','가상대학교','경영정보학과 졸업']],[44,55,75])
    heading(d,'경력')
    table(d,['기간','근무처','담당 업무'],[
      ['2021.03 ~ 2024.02','가상데이터 주식회사','운영 자료 정리와 고객 문의 접수'],
      ['2024.03 ~ 2026.09','예시서비스 주식회사','데이터 품질 점검과 업무 매뉴얼 관리']],[44,55,75])
    heading(d,'업무 역량');para(d,'Excel 자료 정리와 집계, Python을 이용한 반복 업무 처리, SQL 기초 조회')
    heading(d,'자기소개');para(d,'누락된 자료와 중복 기록을 확인하고, 담당자가 이해하기 쉬운 형태로 정리해 왔습니다. 처리 기준을 문서로 남기고 공유 전에 포함 정보를 다시 확인하는 습관을 갖고 있습니다.')
    heading(d,'추천인');para(d,'성명 이지은  직책 팀장\n전화 010-0000-5678  이메일 jieun.lee@example.invalid')
    heading(d,'연락 안내');para(d,'면접 일정은 010-0000-1234 또는 minsu.kim@example.invalid로 안내해 주세요.')
    p=para(d,'2026년 10월 25일  지원자 김민수');p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    save(d,'이력서_검사예시.docx')

    d=base('독서 모임 참가자 안내');para(d,'일시  2026년 11월 7일 토요일 14시부터 16시까지')
    para(d,'장소  '+ADDRESSES[0]);para(d,'대상  사전 신청을 완료한 참가자 8명')
    heading(d,'참가자 명단');table(d,['번호','성명','휴대전화'],[[p['number'],p['name'],p['phone']] for p in PEOPLE[:8]],[18,52,104])
    heading(d,'모임 안내');para(d,'읽고 싶은 책 한 권을 가져와 주세요. 참가 여부가 바뀌면 진행 담당자에게 미리 알려 주세요. 행사 당일 현장에서 출석을 확인합니다.')
    heading(d,'진행 문의');para(d,'담당자 이지은\n전화 010-0000-9011  이메일 event@example.invalid')
    para(d,'안내 작성일  2026년 10월 25일')
    save(d,'행사참가자_안내예시.docx')

    d=base('행사 운영 회의록');para(d,'회의 일시  2026년 10월 25일 10시부터 11시까지')
    para(d,'회의 장소  온라인 회의');para(d,'참석자  김민수, 이지은, 박서윤, 최지우')
    heading(d,'논의 내용');para(d,'1  김민수는 참가 신청 자료의 누락 여부를 확인하고 확정 명단을 정리한다.')
    para(d,'2  이지은은 기념품 수량과 발송 비용을 정리해 다음 회의에서 보고한다.')
    target=para(d,'3  박서윤은 행사 안내문을 작성하고 배포 전 검토를 요청한다.')
    para(d,'4  최지우는 변경 사항을 기록하고 담당자별 완료 여부를 확인한다.')
    heading(d,'배포 일정');table(d,['할 일','기한','전달 방법'],[['안내문 초안 확인','2026.10.28','운영 담당자 검토'],['참가자 안내','2026.10.30','확정 명단 기준 발송']],[64,40,70])
    heading(d,'후속 연락');para(d,'운영 문의 010-0000-9022\n자료 전달 meeting@example.invalid')
    d.add_comment(target.runs,text='공유 전 이지은과 010-0000-9022로 연락해 내용 확인',author='김민수')
    d.core_properties.author='김민수'
    save(d,'회의록_공유예시.docx')

    d=base('경품 배송 요청서');para(d,'발송 예정일  2026년 11월 2일')
    para(d,'요청 부서  예시행사 운영팀\n배송 물품  기념품 세트 5개')
    heading(d,'수취인 명단');table(d,['번호','수취인','연락처','배송 주소'],[[p['number'],p['name'],p['phone'],p['address']] for p in PEOPLE[:5]],[12,24,39,99])
    heading(d,'배송 요청 사항');para(d,'완료 후 주문 번호별 배송 상태를 회신해 주세요. 주소를 확인할 수 없는 건은 발송 전에 운영 담당자에게 문의해 주세요.')
    heading(d,'요청 담당자');para(d,'김민수\n전화 010-0000-9033  이메일 delivery@example.invalid')
    save(d,'배송요청서_공유예시.docx')
    (ROOT/'validation/_build/shared_data.json').write_text(json.dumps({'people':PEOPLE,'addresses':ADDRESSES},ensure_ascii=False,indent=2),encoding='utf-8')
    print('Created 4 synthetic DOCX files')

def create_answer_key():
    needles={**{n:'성명' for n in NAMES},**{p['phone']:'전화번호' for p in PEOPLE},**{p['email']:'이메일' for p in PEOPLE},**{a:'주소' for a in ADDRESSES},
      '010-0000-1234':'전화번호','010-0000-5678':'전화번호','010-0000-9011':'전화번호','010-0000-9022':'전화번호','010-0000-9033':'전화번호',
      'minsu.kim@example.invalid':'이메일','jieun.lee@example.invalid':'이메일','event@example.invalid':'이메일','meeting@example.invalid':'이메일','delivery@example.invalid':'이메일',
      '1998-03-14':'생년월일','28세':'나이'}
    results=[]
    for p in sorted(OUT.glob('*.docx')):
        fields=[]
        with zipfile.ZipFile(p) as z:
            for name in z.namelist():
                if name.endswith(('.xml','.rels')):
                    tree=E.fromstring(z.read(name))
                    for node in tree.iter():
                        if not isinstance(node.tag,str):continue
                        for text in [node.text or '',node.tail or '',*node.attrib.values()]:
                            for needle,label in needles.items():
                                count=text.count(needle)
                                if count:fields.append({'value':needle,'type':label,'count':count,'location':name})
        results.append({'file':p.name,'fields':fields,'checkpoints':sum(x['count'] for x in fields)})
    from openpyxl import load_workbook
    w=load_workbook(OUT/'추첨당첨자_명단예시.xlsx');fields=[]
    for s in w:
        for row in s:
            for c in row:
                text=str(c.value or '')
                for needle,label in needles.items():
                    if needle in text:fields.append({'value':needle,'type':label,'count':text.count(needle),'location':s.title+'!'+c.coordinate,'hidden':s.sheet_state!='visible'})
    results.append({'file':'추첨당첨자_명단예시.xlsx','fields':fields,'checkpoints':sum(x['count'] for x in fields)})
    (OUT/'확인목록.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
    lines=['# 공유 문서 검사 확인 목록','',
      '모든 이름·연락처·주소·기관은 합성 예시입니다. 이 목록은 검사에 넣지 말고 결과와 비교하는 용도로 사용하세요.','',
      '## 사용 순서','',
      '1. 앱의 파일 검사에서 예시 파일 하나를 엽니다.',
      '2. 아래 목록의 값이 해당 위치에서 후보로 표시되는지 확인합니다. 같은 값이 반복된 위치도 각각 확인합니다.',
      '3. 전체 가림 또는 원하는 수정을 선택하고 사본 생성·재검사를 실행합니다.',
      '4. 실제 사본을 열어 선택한 원문이 없어졌는지, 공개 문구·표·서식이 유지되는지 확인합니다.',
      '5. Excel의 숨김 시트와 Word 회의록의 메모·작성자도 확인합니다.','',
      'PDF는 DOCX와 같은 보이는 내용을 가진 검색 가능한 글자층 문서입니다. OCR과 글자층 후보가 겹칠 수 있어 검출 건수를 단순 비교하지 말고 실제 위치를 비교하세요.','',
      '학교·전공·직책·회사·근무기간 등도 이력서 문맥에서 검토할 정보입니다. 아래의 기본 식별·연락 정보 개수와 별도로 확인하세요. 가상 학교·회사 이름이므로 사전 기반 탐지에서 누락될 수 있습니다.','',
      '정답 개수는 알려진 기본 식별·연락 정보의 등장 횟수이며, 기대하는 앱 후보 건수와 같다는 뜻은 아닙니다. 주소 일부만 잡거나 주변 단어까지 잡는 경우도 따로 기록하세요.']
    for result in results:
        lines+=['',f"## {result['file']}",'',f"기본 확인 지점 {result['checkpoints']}곳",'','| 유형 | 값 | 횟수 | 위치 |','|---|---|---:|---|']
        for x in result['fields']:
            loc=x['location']+(' · 숨김 시트' if x.get('hidden') else '')
            lines.append(f"| {x['type']} | {x['value']} | {x['count']} | {loc} |")
    (OUT/'검사_확인목록.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    print([(r['file'],r['checkpoints']) for r in results])
if __name__=='__main__':
    if len(__import__('sys').argv)>1 and __import__('sys').argv[1]=='answers':create_answer_key()
    else:create_documents()
