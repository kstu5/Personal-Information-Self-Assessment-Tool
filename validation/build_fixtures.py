"""Synthetic document regression corpus. No real person's data is used."""
from pathlib import Path
import copy, csv, io, json, zipfile
from urllib.parse import quote
from lxml import etree as E
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

HERE=Path(__file__).resolve().parent; ROOT=HERE.parent; F=HERE/'fixtures'; F.mkdir(exist_ok=True)
CASES=[]
NAME='김민수'; OTHER='이지은'; PHONE='010-0000-1234'; PHONE2='010-0000-5678'; MAIL='sample@example.invalid'; ADDR='서울특별시 마포구 예시로 12, 302호'
STANDARD=f'성명: {NAME}\n전화: {PHONE}\n이메일: {MAIL}\n주소: {ADDR}\n공개 안내: 독서 모임은 토요일 오후에 열립니다.\n'
def check(text,label,count=1,scope=None): return dict(text=text,label=label,count=count,scope=scope)
BASE=[check(NAME,'PERSON'),check(PHONE,'PHONE'),check(MAIL,'EMAIL'),check(ADDR,'ADDRESS')]
def add(file,title,checks=(),**kw): CASES.append(dict(file=file,title=title,checks=list(checks),**kw))
def entries(path):
    with zipfile.ZipFile(path) as z:return {n:z.read(n) for n in z.namelist()}
def savezip(path,data):
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for n,b in data.items():z.writestr(n,b,compress_type=zipfile.ZIP_STORED if n=='mimetype' else zipfile.ZIP_DEFLATED)
def xml(b):return E.fromstring(b)
def xbytes(e):return E.tostring(e,xml_declaration=True,encoding='UTF-8',standalone=True)
def word_base(title):
    d=Document();d.core_properties.author='';d.core_properties.last_modified_by='';d.add_paragraph(title,'Title')
    d.styles['Title'].font.color.rgb=RGBColor(0,0,0)
    for parent in (d.styles['Title'].element,d.paragraphs[0]._p):
        for node in list(parent.iter()):
            if E.QName(node).localname=='pBdr' and node.getparent() is not None:node.getparent().remove(node)
    st=d.styles['Normal'];st.font.name='맑은 고딕';st.font.size=Pt(11)
    st.element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕')
    return d

def main():
    for name,encoding in [('text_utf8.txt','utf-8'),('text_cp949.txt','cp949')]:
        (F/name).write_text(STANDARD,encoding=encoding);add(name,'일반 텍스트와 인코딩',BASE)
    for name,sep in [('table_quoted.csv',','),('table_tabs.tsv','\t')]:
        with (F/name).open('w',encoding='utf-8-sig',newline='') as f:
            csv.writer(f,delimiter=sep).writerows([['항목','값','메모'],['성명',NAME,'쉼표, 줄\n바꿈'],['전화',PHONE,'유지할 공개 문구'],['메일',MAIL,'=1+1']])
        add(name,'구분자와 줄바꿈이 있는 표',BASE[:3],public=['쉼표, 줄\n바꿈','유지할 공개 문구'])
    (F/'text_variants.txt').write_text(f'성명: 김 민 수\n전화: 010 0000 1234\n전화: 010.0000.5678\n이메일: sample @ example.invalid\n주소: 서울특별시 마포구 예시로 12\n상세주소: 302호\n',encoding='utf-8')
    add('text_variants.txt','띄어쓰기와 구분자 변형',[check('김 민 수','PERSON'),check('010 0000 1234','PHONE'),check('010.0000.5678','PHONE'),check('sample @ example.invalid','EMAIL'),check('서울특별시 마포구 예시로 12','ADDRESS'),check('302호','ADDRESS')],category='challenge')
    (F/'text_negative.txt').write_text('자료 준비 안내\n책상 아래에 두었던 책을 함께 읽습니다.\n총 12권을 준비하고 토요일 오후 2시에 만나요.\n가을 독서 모임에는 누구나 참여할 수 있습니다.\n',encoding='utf-8')
    add('text_negative.txt','개인정보 없는 대조 문서',[],negative=True)

    d=word_base('참가 신청서');p=d.add_paragraph('담당자 ');p.add_run('김').bold=True;p.add_run('민').italic=True;p.add_run('수')
    d.add_paragraph(f'전화 {PHONE} 이메일 {MAIL}');d.add_paragraph(f'주소 {ADDR}')
    t=d.add_table(rows=2,cols=2);t.style='Table Grid';t.cell(0,0).text='공동 담당';t.cell(0,1).text=OTHER;t.cell(1,0).text='연락';t.cell(1,1).text=PHONE2
    d.sections[0].header.paragraphs[0].text=f'담당 {NAME}';d.sections[0].footer.paragraphs[0].text=f'문의 {PHONE}'
    d.add_comment(p.runs,text=f'검토 담당 {OTHER}',author=NAME)
    deletion=OxmlElement('w:del');deletion.set(qn('w:author'),OTHER);rr=OxmlElement('w:r');tt=OxmlElement('w:delText');tt.text=f'이전 연락 {PHONE2}';rr.append(tt);deletion.append(rr);p._p.append(deletion)
    d.core_properties.author=NAME;d.core_properties.subject=f'검토 {OTHER}';d.save(F/'word_structures.docx')
    add('word_structures.docx','분리 서식 본문 표 머리말 꼬리말 메모 변경기록 속성',[check(NAME,'PERSON',4),check(PHONE,'PHONE',2),check(MAIL,'EMAIL'),check(ADDR,'ADDRESS'),check(OTHER,'PERSON',4),check(PHONE2,'PHONE',2)])

    d=word_base('선택 수정 예시');d.add_paragraph(f'😀 담당 {NAME} / 공동 담당 {NAME}');d.add_paragraph(f'연락 {PHONE}');d.add_paragraph('공개 문구는 그대로 남아야 합니다.');d.save(F/'word_repeated.docx')
    add('word_repeated.docx','이모지 반복 위치 선택 수정',[check(NAME,'PERSON',2),check(PHONE,'PHONE')],actions=True,public=['공개 문구는 그대로 남아야 합니다.'])

    d=word_base('줄바꿈 입력 예시');p=d.add_paragraph('성명: 김');p.add_run().add_break();p.add_run('민수');p=d.add_paragraph('전화: 010-');p.add_run().add_break();p.add_run('0000-1234');d.save(F/'word_linebreak.docx')
    add('word_linebreak.docx','문단 내부 강제 줄바꿈',[check(NAME,'PERSON'),check(PHONE,'PHONE')],category='challenge',allow_line_breaks=True)

    d=word_base('링크가 있는 연락 안내');p=d.add_paragraph('문의 ')
    hl=OxmlElement('w:hyperlink');hl.set(qn('r:id'),d.part.relate_to('mailto:'+MAIL,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=MAIL;r.append(t);hl.append(r);p._p.append(hl)
    d.add_paragraph(f'담당 {NAME}');d.save(F/'word_link.docx');add('word_link.docx','표시 글자와 실제 링크 주소',[check(MAIL,'EMAIL',2),check(NAME,'PERSON')])

    d=word_base('부가 정보가 있는 문서');d.add_paragraph('공개 안내만 있는 본문입니다.')
    # A normal hyperlink's user-facing tooltip can contain personal information.
    p=d.add_paragraph();hl=OxmlElement('w:hyperlink');hl.set(qn('r:id'),d.part.relate_to('https://example.invalid',RT.HYPERLINK,is_external=True));hl.set(qn('w:tooltip'),f'연락 {PHONE}');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='행사 안내';r.append(t);hl.append(r);p._p.append(hl)
    variables=OxmlElement('w:docVars');v=OxmlElement('w:docVar');v.set(qn('w:name'),'담당자');v.set(qn('w:val'),NAME);variables.append(v);d.settings.element.append(variables)
    # VML textbox without an image. Alternative text and shape name are real attributes.
    p=d.add_paragraph();pic=OxmlElement('w:pict');shape=E.SubElement(pic,'{urn:schemas-microsoft-com:vml}rect',id='ContactShape',style='width:180pt;height:30pt');shape.set('alt',f'검토자 {OTHER}')
    tb=E.SubElement(shape,'{urn:schemas-microsoft-com:vml}textbox');tc=OxmlElement('w:txbxContent');wp=OxmlElement('w:p');wr=OxmlElement('w:r');wt=OxmlElement('w:t');wt.text='공개 텍스트 상자';wr.append(wt);wp.append(wr);tc.append(wp);tb.append(tc);p._p.append(pic)
    d.save(F/'word_attributes.docx');add('word_attributes.docx','툴팁 문서 변수 대체 텍스트',[check(PHONE,'PHONE'),check(NAME,'PERSON'),check(OTHER,'PERSON')],category='coverage')

    d=word_base('각주 연락 안내');d.add_paragraph('본문은 공개 안내입니다.');d.save(F/'word_footnotes.docx');data=entries(F/'word_footnotes.docx')
    ns='http://schemas.openxmlformats.org/wordprocessingml/2006/main';foot=E.Element('{'+ns+'}footnotes',nsmap={'w':ns});n=E.SubElement(foot,'{'+ns+'}footnote',{'{'+ns+'}id':'1'});p=E.SubElement(n,'{'+ns+'}p');r=E.SubElement(p,'{'+ns+'}r');E.SubElement(r,'{'+ns+'}t').text=f'각주 담당 {NAME} 연락 {PHONE}'
    data['word/footnotes.xml']=xbytes(foot);rel=xml(data['word/_rels/document.xml.rels']);E.SubElement(rel,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',Id='rIdFootQA',Type=RT.FOOTNOTES,Target='footnotes.xml');data['word/_rels/document.xml.rels']=xbytes(rel)
    ct=xml(data['[Content_Types].xml']);E.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Override',PartName='/word/footnotes.xml',ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml');data['[Content_Types].xml']=xbytes(ct)
    body=xml(data['word/document.xml']);p=next(body.iter('{'+ns+'}p'));r=E.SubElement(p,'{'+ns+'}r');E.SubElement(r,'{'+ns+'}footnoteReference',{'{'+ns+'}id':'1'});data['word/document.xml']=xbytes(body);savezip(F/'word_footnotes.docx',data)
    add('word_footnotes.docx','각주에만 있는 개인정보',[check(NAME,'PERSON'),check(PHONE,'PHONE')])

    # Modify OOXML flags for deterministic hidden/metadata cases. No application code is used.
    path=F/'xlsx_values.xlsx';data=entries(path);tree=xml(data['xl/workbook.xml']);sheets=[e for e in tree.iter() if E.QName(e).localname=='sheet'];sheets[1].set('state','veryHidden');data['xl/workbook.xml']=xbytes(tree)
    sheet=xml(data['xl/worksheets/sheet1.xml']);
    for row in sheet.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}row'):
        if row.get('r')=='3':row.set('hidden','1')
    data['xl/worksheets/sheet1.xml']=xbytes(sheet);savezip(path,data)
    add(path.name,'숨김 시트 숨김 행 일반 메모',BASE[:3]+[check(OTHER,'PERSON',2),check(PHONE2,'PHONE',2)])
    add('xlsx_formulas.xlsx','이름 전화번호 파생 수식과 값 고정',[check(NAME,'PERSON',2),check(PHONE,'PHONE'),check(PHONE2,'PHONE'),check(MAIL,'EMAIL')])
    add('xlsx_blocked.xlsx','미지원 유효성 검사 차단',[],expected_block=True)
    # Deliberate sheet-name collision caused by applying one mask to distinct names.
    data=entries(F/'xlsx_values.xlsx');tree=xml(data['xl/workbook.xml']);sheets=[e for e in tree.iter() if E.QName(e).localname=='sheet'];sheets[0].set('name',NAME);sheets[1].set('name',OTHER);data['xl/workbook.xml']=xbytes(tree);savezip(F/'xlsx_sheet_names.xlsx',data)
    add('xlsx_sheet_names.xlsx','개인 이름 시트의 일괄 전체 가림',[check(NAME,'PERSON',2),check(PHONE,'PHONE'),check(MAIL,'EMAIL'),check(OTHER,'PERSON',3),check(PHONE2,'PHONE',2)],category='challenge')

    data=entries(F/'pptx_body.pptx');tree=xml(data['ppt/slides/slide2.xml']);tree.set('show','0');data['ppt/slides/slide2.xml']=xbytes(tree);savezip(F/'pptx_body.pptx',data)
    add('pptx_body.pptx','슬라이드 본문 숨김 슬라이드 발표자 노트',[check(NAME,'PERSON'),check(PHONE,'PHONE'),check(MAIL,'EMAIL'),check(OTHER,'PERSON',2),check(PHONE2,'PHONE',2)])
    data=entries(F/'pptx_body.pptx');tree=xml(data['ppt/slides/slide1.xml']);node=next(e for e in tree.iter() if E.QName(e).localname=='cNvPr');node.set('descr','추가 연락 010-0000-9999');node.set('name','담당 박서윤');data['ppt/slides/slide1.xml']=xbytes(tree);savezip(F/'pptx_attributes.pptx',data)
    add('pptx_attributes.pptx','도형 대체 텍스트와 이름',[check('010-0000-9999','PHONE'),check('박서윤','PERSON')],category='coverage')

    # Reuse a local synthetic HWPX template solely to retain valid Hancom style/layout parts.
    template=HERE/'templates/hwpx_synthetic_base.hwpx';data=entries(template)
    for name,raw in list(data.items()):
        if name.endswith(('.xml','.hpf')):
            try:
                tree=xml(raw)
                for e in tree.iter():
                    if not isinstance(e.tag,str):continue
                    if E.QName(e).localname in {'t','text'} and e.text and ('김민수' in e.text or '이지은' in e.text or '010-' in e.text):e.text=e.text.replace('김민수',NAME).replace('이지은',OTHER)
                data[name]=xbytes(tree)
            except E.XMLSyntaxError:pass
    savezip(F/'hancom_notes.hwpx',data)
    add('hancom_notes.hwpx','HWPX 본문 메모 작성자와 미리보기 제거',[check(NAME,'PERSON',2),check(OTHER,'PERSON',3),check(PHONE,'PHONE',3)],template_source=str(template.relative_to(ROOT)))

    font=Path('C:/Windows/Fonts/malgun.ttf');pdfmetrics.registerFont(TTFont('QA-Malgun',str(font)))
    path=F/'pdf_text.pdf';c=canvas.Canvas(str(path),pagesize=(595,842));c.setTitle('문서 개인정보 검토 예시');c.setAuthor(NAME);c.setFont('QA-Malgun',14)
    for i,line in enumerate(STANDARD.splitlines()):c.drawString(45,785-i*30,line)
    c.save();add(path.name,'검색 가능한 PDF 글자층',BASE)
    path=F/'pdf_annotation.pdf';c=canvas.Canvas(str(path),pagesize=(595,842));c.setFont('QA-Malgun',14);c.drawString(45,785,'공개 안내 문서');c.save()
    import pymupdf as fitz
    doc=fitz.open(path);doc[0].add_text_annot((50,100),f'담당 {NAME} 연락 {PHONE}');tmp=path.with_suffix('.tmp.pdf');doc.save(tmp);doc.close();tmp.replace(path)
    add(path.name,'PDF 주석에만 있는 개인정보',[check(NAME,'PERSON'),check(PHONE,'PHONE')],expected_block=True,category='coverage')

    (HERE/'manifest.json').write_text(json.dumps({'data_policy':'All identities and contact details are synthetic; example.invalid is a reserved domain. No real person is represented.','cases':CASES},ensure_ascii=False,indent=2),encoding='utf-8')
    print(f'Created {len(CASES)} document fixtures',flush=True)

if __name__=='__main__': main()
