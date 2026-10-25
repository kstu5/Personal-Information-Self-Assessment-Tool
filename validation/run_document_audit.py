"""Exercise the real authenticated API and actual pinned A model, never a mock detector."""
from __future__ import annotations
from pathlib import Path
import collections, hashlib, html, importlib.metadata, json, platform, re, subprocess, sys, time, traceback, zipfile
from urllib.parse import unquote
ROOT=Path(__file__).resolve().parents[1];HERE=ROOT/'validation';sys.path.insert(0,str(ROOT))
from lxml import etree as E
from fastapi.testclient import TestClient
from backend.app import create_app

def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def u16(s):return len(s.encode('utf-16-le'))//2
def write(path,data):path.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
def external_audit(path,needles):
    """Independently inspect saved content, including XML attributes not exposed by the app."""
    pieces=[]
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as z:
            for n in z.namelist():
                if n.endswith(('.xml','.rels','.hpf','.txt')):
                    raw=z.read(n)
                    try:text=raw.decode('utf-8-sig')
                    except UnicodeError:continue
                    pieces.append((n+' raw',unquote(html.unescape(text))))
                    if n.endswith(('.xml','.rels','.hpf')):
                        try:
                            tree=E.fromstring(raw)
                            pieces.append((n+' joined text',''.join(tree.itertext())))
                            for e in tree.iter():
                                if isinstance(e.tag,str):
                                    for k,v in e.attrib.items():pieces.append((n+' @'+E.QName(k).localname,unquote(v)))
                        except E.XMLSyntaxError:pass
    elif path.suffix=='.pdf':
        import pymupdf as fitz
        with fitz.open(path) as d:
            pieces.append(('PDF metadata',str(d.metadata)))
            for i,p in enumerate(d):
                pieces.append((f'page {i+1} native text',p.get_text()))
                for a in p.annots() or []:pieces.append((f'page {i+1} annotation',str(a.info)))
            for i in range(1,d.xref_length()):
                pieces.append((f'xref {i}',d.xref_object(i)))
                if d.xref_is_stream(i):
                    stream=d.xref_stream(i)
                    for encoding in ('utf-8','utf-16-be'):
                        pieces.append((f'xref {i} stream {encoding}',stream.decode(encoding,errors='ignore')))
    else:
        raw=path.read_bytes()
        for enc in ('utf-8-sig','cp949','utf-16'):
            try:pieces=[('text',raw.decode(enc))];break
            except UnicodeError:continue
    hits={}
    for needle in needles:
        where=sorted(set(n for n,text in pieces if needle in text))
        if where:hits[needle]=where
    return {'remaining':hits,'publicTexts':[text for n,text in pieces if n.endswith('joined text') or n=='text' or 'native text' in n]}

def checkpoints(case,job):
    parts=(job.get('manifest') or {}).get('parts',[]);cs=job.get('candidates',[]);out=[]
    for gt in case['checks']:
        locations=[];full=[];partial=[];overextended=[]
        for p in parts:
            if gt.get('scope') and gt['scope'] not in p['location']:continue
            # PDF OCR duplicates the underlying native text. Prefer native text in scoring.
            if case['file'].endswith('.pdf') and 'PDF 글자층' not in p['location']:continue
            projection=[i for i,ch in enumerate(p['text']) if not case.get('allow_line_breaks') or ch not in '\r\n']
            searchable=''.join(p['text'][i] for i in projection)
            for match in re.finditer(re.escape(gt['text']),searchable):
                left,right=projection[match.start()],projection[match.end()-1]+1
                a,b=u16(p['text'][:left]),u16(p['text'][:right]);locations.append(p['location'])
                overlap=[c for c in cs if c['part']==p['id'] and c['start']<b and c['end']>a]
                covered=[c for c in overlap if c['label']==gt['label'] and c['start']<=a and c['end']>=b]
                example={'location':p['location'],'expected':gt['text'],'candidates':[{'text':c['text'],'label':c['label'],'start':c['start'],'end':c['end']} for c in overlap]}
                if covered:
                    full.append(example)
                    if any(c['start']<a or c['end']>b for c in covered):overextended.append(example)
                elif overlap:partial.append(example)
        count=gt['count'];nfull=min(count,len(full));npart=min(count-nfull,len(partial));read=min(count,len(locations))
        out.append({**gt,'extracted':read,'fullyDetected':nfull,'partiallyDetected':npart,'notExtracted':count-read,'notFullyDetected':count-nfull,'locations':locations,'fullExamples':full,'partialExamples':partial,'overextendedExamples':overextended})
    return out

def main():
    import argparse
    parser=argparse.ArgumentParser();parser.add_argument('--output-dir',type=Path)
    args=parser.parse_args()
    output=args.output_dir.resolve() if args.output_dir else HERE
    if not output.is_relative_to(HERE):raise ValueError('Audit output must stay within the test directory')
    manifest=json.loads((HERE/'manifest.json').read_text('utf-8'));results=output/'results';results.mkdir(parents=True,exist_ok=True);reviewed=output/'reviewed';reviewed.mkdir(parents=True,exist_ok=True)
    commit=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip()
    env={'commit':commit,'python':sys.version,'platform':platform.platform(),'testedAt':'2026-10-25 Asia/Seoul','packages':{}}
    for p in ('torch','transformers','ko-pii','lxml','openpyxl','tesserocr','PyMuPDF','fastapi','python-docx'):
        try:env['packages'][p]=importlib.metadata.version(p)
        except importlib.metadata.PackageNotFoundError:env['packages'][p]='not found'
    write(results/'environment.json',env)
    app=create_app(19264);report=[]
    with TestClient(app,base_url='http://127.0.0.1:19264') as client:
        client.headers['X-App-Token']=app.state.token
        def post(route,data):
            r=client.post('/api/documents/'+route,json=data)
            if r.status_code!=200:raise RuntimeError(f'{route}: {r.status_code} {r.text}')
            return r.json()
        def wait(ident,limit=300):
            end=time.monotonic()+limit;last=None
            while time.monotonic()<end:
                j=post('status',{'id':ident})
                if j['state'] in {'ready','error','cancelled'}:return j
                if j['progress']!=last:print('  '+j['progress'],flush=True);last=j['progress']
                time.sleep(.15)
            raise TimeoutError('document job exceeded audit timeout')
        def export(ident,rev,choices,mode,case,edits=None):
            started=time.monotonic();post('export',{'id':ident,'revision':rev,'choices':choices,'regions':[],'textEdits':edits or {}});done=wait(ident)
            write(results/(Path(case['file']).stem+'_'+mode+'_job.json'),done)
            row={'saved':False,'seconds':round(time.monotonic()-started,3),'error':done.get('error')}
            if done.get('exportResult'):
                r=client.get(f'/api/documents/{ident}/download');r.raise_for_status();suffix=done['manifest']['outputSuffix'];out=reviewed/(Path(case['file']).stem+'__'+mode+suffix);out.write_bytes(r.content)
                assert sha(out)==done['exportResult']['sha256']
                audit=external_audit(out,[x['text'] for x in case['checks']]);write(results/(out.stem+'_independent.json'),audit)
                row.update(saved=True,file=str(out.relative_to(HERE)),sha256=sha(out),remaining=audit['remaining'],rescanCandidates=done['exportResult']['candidates'],rescanCount=len(done['exportResult']['candidates']))
                row['publicPreserved']=all(any(t in s for s in audit['publicTexts']) for t in case.get('public',[]))
            return row
        for case in manifest['cases']:
            name=case['file'];src=HERE/'fixtures'/name;before=sha(src);started=time.monotonic();ident=None
            row={'file':name,'sourceSHA256':before,'title':case['title'],'category':case.get('category','normal'),'expectedBlock':case.get('expected_block',False)}
            print('START '+name,flush=True)
            try:
                r=client.post('/api/documents/upload',params={'filename':name},content=src.read_bytes(),headers={'Content-Type':'application/octet-stream'});r.raise_for_status();ident=r.json()['id'];j=wait(ident)
                write(results/(src.stem+'_input_job.json'),j)
                row.update(state=j['state'],error=j.get('error'),readSeconds=round(time.monotonic()-started,3),candidateCount=len(j.get('candidates',[])),checkpoints=checkpoints(case,j))
                if j.get('manifest'):
                    row.update(blockers=j['manifest']['blockers'],warnings=j['manifest']['warnings'],parts=len(j['manifest']['parts']),assets=len(j['manifest']['assets']),calculation=j['manifest']['calculation'])
                blocked=j['state']=='error' or bool(row.get('blockers'))
                row['expectedBlockObserved']=row['expectedBlock'] and blocked
                if row['expectedBlock']:
                    # Ensure the actual export route also refuses incomplete coverage.
                    refused=client.post('/api/documents/export',json={'id':ident,'revision':j.get('revision',''),'choices':{},'regions':[]})
                    row['exportRefusalStatus']=refused.status_code
                elif j['state']=='ready' and not blocked:
                    row['autoFull']=export(ident,j['revision'],{c['id']:{'mode':'full'} for c in j['candidates']},'auto_full',case)
                    if case['checks'] and not name.endswith('.pdf'):
                        edits={}
                        for p in j['manifest']['parts']:
                            if p.get('asset'):continue
                            value=p['text']
                            for check in case['checks']:value=value.replace(check['text'],'[삭제]')
                            if value!=p['text']:edits[p['id']]=value
                        row['manualRepair']=export(ident,j['revision'],{},'manual_repair',case,edits)
                    if case.get('actions'):
                        names=[c for c in j['candidates'] if c['label']=='PERSON' and c['text']=='김민수'];phones=[c for c in j['candidates'] if c['label']=='PHONE']
                        choices={names[0]['id']:{'mode':'partial'}} if names else {}
                        choices.update({c['id']:{'mode':'full'} for c in phones})
                        row['selective']=export(ident,j['revision'],choices,'selective',case)
                        if row['selective']['saved']:
                            from docx import Document
                            text='\n'.join(p.text for p in Document(HERE/row['selective']['file']).paragraphs)
                            row['selective']['assertions']={'oneUnselectedNamePreserved':text.count('김민수')==1,'partialNameApplied':text.count('김**')==1,'phoneRemoved':'010-0000-1234' not in text,'publicPreserved':'공개 문구는 그대로 남아야 합니다.' in text}
                        row['custom']=export(ident,j['revision'],{c['id']:{'mode':'custom','text':'[담당자]'} for c in names},'custom',case)
                        if row['custom']['saved']:
                            from docx import Document
                            text='\n'.join(p.text for p in Document(HERE/row['custom']['file']).paragraphs)
                            row['custom']['assertions']={'bothNamesReplaced':text.count('[담당자]')==2 and '김민수' not in text,'unselectedPhonePreserved':'010-0000-1234' in text,'publicPreserved':'공개 문구는 그대로 남아야 합니다.' in text}
                if case.get('negative'):row['falsePositiveCandidates']=j.get('candidates',[])
            except Exception as e:
                row['harnessError']=str(e);row['traceback']=traceback.format_exc()
            finally:
                row['originalUnchanged']=sha(src)==before
                if ident:
                    try:post('remove',{'id':ident})
                    except Exception as e:row['cleanupError']=str(e)
                report.append(row);write(results/'audit.json',{'environment':env,'cases':report});print('DONE '+name+' '+json.dumps({k:row.get(k) for k in ('state','error','harnessError','candidateCount')},ensure_ascii=False),flush=True)
    normal=[r for r in report if not r['expectedBlock']];allchecks=[c for r in normal for c in r.get('checkpoints',[])]
    summary={'files':len(report),'expectedBlockedFiles':sum(r['expectedBlock'] for r in report),'expectedBlocksObserved':sum(r.get('expectedBlockObserved',False) for r in report),'readyFiles':sum(r.get('state')=='ready' for r in normal),'exportedFiles':sum(r.get('autoFull',{}).get('saved',False) for r in normal),'originalsUnchanged':sum(r['originalUnchanged'] for r in report),'targetCheckpoints':sum(c['count'] for c in allchecks),'fullyDetectedCheckpoints':sum(c['fullyDetected'] for c in allchecks),'partiallyDetectedCheckpoints':sum(c['partiallyDetected'] for c in allchecks),'notExtractedCheckpoints':sum(c['notExtracted'] for c in allchecks),'filesWithIndependentResiduals':[r['file'] for r in normal if r.get('autoFull',{}).get('remaining')],'zeroRescanButResiduals':[r['file'] for r in normal if r.get('autoFull',{}).get('saved') and r['autoFull']['rescanCount']==0 and r['autoFull']['remaining']],'harnessErrors':[r['file'] for r in report if r.get('harnessError')]}
    summary['overextendedExamples']=sum(len(c['overextendedExamples']) for c in allchecks)
    write(results/'summary.json',summary);print(json.dumps(summary,ensure_ascii=False,indent=2),flush=True)

if __name__=='__main__':main()
