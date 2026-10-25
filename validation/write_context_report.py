"""Check local evidence and write a concise, shareable context-audit summary."""
from pathlib import Path
import hashlib
import json
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / 'validation')]
from run_document_audit import external_audit

PUBLIC_DATES = {
    '이력서_검사예시': ['2026년 10월 25일'],
    '행사참가자_안내예시': ['2026년 11월 7일', '2026년 10월 25일'],
    '배송요청서_공유예시': ['2026년 11월 2일'],
    '회의록_공유예시': ['2026년 10월 25일', '2026.10.28', '2026.10.30'],
}

def main():
    raw = ROOT / 'local-only/context-review'
    before = json.loads((raw / 'before/results.json').read_text('utf-8'))
    after = json.loads((raw / 'after/results.json').read_text('utf-8'))
    probes = json.loads((raw / 'probes.json').read_text('utf-8'))
    regression = json.loads((ROOT / 'validation/qa/context-regression/results/summary.json').read_text('utf-8'))
    for name, digest in after['sourceSHA256'].items():
        assert hashlib.sha256((ROOT / name).read_bytes()).hexdigest() == digest, 'Source changed since audit: ' + name
    assert regression['fullyDetectedCheckpoints'] == regression['targetCheckpoints'] == 89
    assert not regression['harnessErrors'] and not regression['filesWithIndependentResiduals']
    assert all(not row['after']['missing'] and not row['after']['unwanted'] for row in probes)
    previous = {row['file']: row for row in before['files']}
    rows = []
    for row in after['files']:
        old = previous[row['file']]
        source = ROOT / 'examples/documents' / row['file']
        assert row['sha256'] == old['sha256'] == hashlib.sha256(source.read_bytes()).hexdigest()
        assert row['full'] == row['expected'] and row['originalUnchanged'] and row['exported']
        assert not row['residuals'] and not row['rescan']
        target = raw / 'after' / (source.stem + '__reviewed' + source.suffix)
        # Re-open the actual exported file, independently of the app rescan.
        independent = external_audit(target, {check['text'] for check in row['checks']})
        assert not independent['remaining']
        if source.suffix == '.xlsx':
            from openpyxl import load_workbook
            a, b = load_workbook(source), load_workbook(target)
            assert b['내부_배송정보'].sheet_state == 'hidden'
            assert all(a['공유용_명단'].cell(r, 5).value == b['공유용_명단'].cell(r, 5).value for r in range(5, 15))
            a.close(); b.close()
            preserved = {'publicationDateCells': 10}
        else:
            values = PUBLIC_DATES[source.stem]
            assert all(any(value in text for text in independent['publicTexts']) for value in values)
            preserved = {'dates': values}
        rows.append({'file': row['file'], 'sourceSHA256': row['sha256'], 'expected': row['expected'],
                     'beforeFull': old['full'], 'afterFull': row['full'], 'beforePartial': old['partial'],
                     'afterPartial': row['partial'], 'beforeCandidates': old['candidates'], 'afterCandidates': row['candidates'],
                     'beforeExtras': old['extras'], 'afterExtras': row['extras'], 'exported': True,
                     'knownOriginalsRemaining': [], 'originalUnchanged': True, 'preservedPublicDates': preserved})
    summary = {'baselineCommit': before['commit'], 'afterSourceSHA256': after['sourceSHA256'],
               'expected': sum(r['expected'] for r in rows), 'beforeFull': sum(r['beforeFull'] for r in rows),
               'afterFull': sum(r['afterFull'] for r in rows), 'files': rows,
               'probes': probes, 'previousCorpus': regression,
               'limits': 'Synthetic checkpoints only; PDF scoring uses native text; no general accuracy or complete removal guarantee.'}
    folder = ROOT / 'validation/context-review'; folder.mkdir(exist_ok=True)
    (folder / 'summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: summary[k] for k in ('expected', 'beforeFull', 'afterFull')}, ensure_ascii=False))

if __name__ == '__main__':
    main()
