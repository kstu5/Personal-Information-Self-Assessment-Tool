from pathlib import Path
import hashlib
import io
import json
import re
import sys
import time
import zipfile

import pytest
from fastapi.testclient import TestClient
from backend.app import create_app
from backend.documents import DocumentError
from backend.file_review import FileReviewStore, utf16

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'.runtime'/'file-python'))
from backend.file_formats import Document, FormatError, read_zip


class RegexDetector:
    def check(self,title,body):
        found=[]
        for m in re.finditer(r'김민수|이지은|010-\d{4}-\d{4}|sample@example\.invalid',body):
            label='PHONE' if m[0].startswith('010') else 'EMAIL' if '@' in m[0] else 'PERSON'
            found.append({'id':str(m.start()),'source':'body','start':utf16(body[:m.start()]),'end':utf16(body[:m.end()]),'text':m[0],'label':label,'labelName':label,'risk':'HIGH'})
        return {'candidates':found,'engine':'test-only'}


def wait(store,ident,timeout=200):
    end=time.monotonic()+timeout
    while time.monotonic()<end:
        result=store.status(ident)
        if result['state'] in {'ready','error','cancelled'}:return result
        time.sleep(.05)
    pytest.fail('file operation did not finish')


@pytest.fixture
def store():
    s=FileReviewStore(RegexDetector())
    yield s
    root=s.root;s.close();assert not root.exists()


def test_copy_only_unicode_occurrences_export_and_stale_source(store,tmp_path):
    original=tmp_path/'모임.txt';original.write_text('😀 이름 김민수\n전화 010-0000-1234\n김민수는 박하차를 준비합니다.',encoding='utf-8')
    before=original.read_bytes();job=store.upload(original.name,before,original)
    ready=wait(store,job['id']);assert ready['state']=='ready',ready
    names=[c for c in ready['candidates'] if c['text']=='김민수'];assert len(names)==2
    choices={names[0]['id']:{'mode':'partial'}}
    phone=next(c for c in ready['candidates'] if c['label']=='PHONE');choices[phone['id']]={'mode':'full'}
    with pytest.raises(DocumentError):store.export(job['id'],'stale',choices,[])
    store.export(job['id'],ready['revision'],choices,[])
    done=wait(store,job['id']);assert done['exportResult'],done.get('error')
    path,_=store.output(job['id']);text=path.read_text('utf-8')
    assert '김**' in text and text.count('김민수')==1 and '010-0000-1234' not in text
    assert original.read_bytes()==before
    with pytest.raises(DocumentError):store.save_new(job['id'],original)
    output=tmp_path/'새 사본.txt';store.save_new(job['id'],output);assert output.read_bytes()==path.read_bytes()
    original.write_text('changed',encoding='utf-8')
    with pytest.raises(DocumentError):store.output(job['id'])


def test_api_file_auth_binary_validation_and_independent_sessions(tmp_path):
    app=create_app(19002)
    with TestClient(app,base_url='http://127.0.0.1:19002') as c:
        assert c.post('/api/documents/upload?filename=a.txt',content=b'x',headers={'Content-Type':'application/octet-stream'}).status_code==403
        assert c.post('/api/documents/list',json={}).status_code==403
        c.headers['X-App-Token']=app.state.token;app.state.documents.detector=RegexDetector()
        assert c.post('/api/documents/upload?filename=../x.txt',content=b'x',headers={'Content-Type':'application/octet-stream'}).status_code==400
        assert c.post('/api/documents/upload?filename=a.exe',content=b'x',headers={'Content-Type':'application/octet-stream'}).status_code==400
        assert c.post('/api/documents/upload?filename=a.txt',content=b'x',headers={'Content-Type':'application/json'}).status_code==415
        r=c.post('/api/documents/upload?filename=a.txt',content='김민수'.encode(),headers={'Content-Type':'application/octet-stream'})
        assert r.status_code==200,r.text
        job=wait(app.state.documents,r.json()['id']);assert job['state']=='ready'
        restored=c.post('/api/documents/list',json={}).json()['files']
        assert len(restored)==1 and restored[0]['id']==job['id']
        assert 'work' not in restored[0] and 'original' not in restored[0]
        result=c.post('/api/documents/status',json={'id':job['id']});assert result.status_code==200
        assert 'target' not in result.text and str(tmp_path) not in result.text
        assert c.get(f"/api/documents/{job['id']}/download").status_code==400
        assert c.post('/api/documents/export',json={'id':job['id'],'revision':job['revision'],'choices':{'unknown':{'mode':'full'}},'regions':[]}).status_code==400
        assert c.post('/api/documents/remove',json={'id':job['id']}).status_code==200
        assert c.post('/api/documents/list',json={}).json()['files']==[]
        assert c.post('/api/documents/status',json={'id':job['id']}).status_code==400


def test_zip_traversal_duplicates_and_entities_rejected(tmp_path):
    from backend.file_formats import parse_xml
    p=tmp_path/'bad.docx'
    with zipfile.ZipFile(p,'w') as z:z.writestr('../leak','x')
    with pytest.raises(FormatError):read_zip(p)
    with pytest.raises(FormatError):parse_xml(b'<!DOCTYPE x [<!ENTITY z SYSTEM "file:///secret">]><x>&z;</x>')


def test_docx_comments_revisions_run_boundaries_and_thumbnail(tmp_path):
    from docx import Document as Word
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    p=tmp_path/'source.docx';d=Word();para=d.add_paragraph();r=para.add_run('담당 김');r.bold=True;para.add_run('민수 입니다.');d.add_comment(para.runs,text='메모 이지은',author='김민수')
    deletion=OxmlElement('w:del');deletion.set(qn('w:author'),'김민수');run=OxmlElement('w:r');text=OxmlElement('w:delText');text.text='이전 이지은';run.append(text);deletion.append(run);para._p.append(deletion)
    d.core_properties.author='김민수';d.save(p)
    doc=Document(p,tmp_path).load();assert not doc.blockers
    text='\n'.join(p['text'] for p in doc.parts);assert '메모 이지은' in text and '이전 이지은' in text
    edits={p['id']:p['text'].replace('김민수','김**').replace('이지은','이**') for p in doc.parts if '김민수' in p['text'] or '이지은' in p['text']}
    output=tmp_path/'copy.docx';doc.export(output,edits,[])
    assert Word(output).paragraphs[0].runs[0].bold
    with zipfile.ZipFile(output) as z:
        content=b'\n'.join(z.read(n) for n in z.namelist() if n.endswith('.xml'))
        assert '김민수'.encode() not in content and '이지은'.encode() not in content
        assert not any(n.lower().startswith('docprops/thumbnail') for n in z.namelist())
    # A direct custom insertion at the end of a run must not fail or lose surrounding text.
    simple=Word();simple.add_paragraph('김민수');simple.save(p);doc=Document(p,tmp_path).load();part=next(p for p in doc.parts if p['text']=='김민수')
    doc.export(tmp_path/'append.docx',{part['id']:'김민수 담당자'},[])
    assert Word(tmp_path/'append.docx').paragraphs[0].text=='김민수 담당자'


def test_excel_external_formula_blocked_before_calculation(tmp_path,monkeypatch):
    from openpyxl import Workbook
    p=tmp_path/'source.xlsx';w=Workbook();w.active['A1']='=WEBSERVICE("https://example.invalid")';w.save(p)
    def fail(*args):pytest.fail('unsafe formula reached calculator')
    monkeypatch.setattr(Document,'recalculate',fail)
    with pytest.raises(FormatError):Document(p,tmp_path).load()


def test_excel_formatted_numbers_hidden_comments_and_literal_custom(tmp_path):
    from openpyxl import Workbook,load_workbook
    from openpyxl.comments import Comment
    p=tmp_path/'source.xlsx';w=Workbook();s=w.active;s['A1']=1000001234;s['A1'].number_format='000-0000-0000';s['A1'].comment=Comment('김민수','이지은');h=w.create_sheet('숨김');h.sheet_state='hidden';h['A1']='김민수';w.save(p)
    doc=Document(p,tmp_path).load();assert not doc.blockers
    parts=doc.parts;assert any(p['text']=='010-0000-1234' for p in parts)
    assert any('숨김 시트' in p['location'] for p in parts)
    assert any('작성자' in p['location'] and p['text']=='이지은' for p in parts)
    target=next(p for p in parts if p['text']=='010-0000-1234');out=tmp_path/'copy.xlsx';doc.export(out,{target['id']:'=1+1'},[])
    assert load_workbook(out).active['A1'].data_type=='s'


def test_image_opaque_region_metadata_and_source_hash(store,tmp_path):
    from PIL import Image,PngImagePlugin
    p=tmp_path/'image.png';im=Image.new('RGB',(100,80),'white');meta=PngImagePlugin.PngInfo();meta.add_text('private','김민수');im.save(p,pnginfo=meta)
    raw=p.read_bytes();j=store.upload(p.name,raw,p);r=wait(store,j['id']);assert r['state']=='ready',r
    store.export(j['id'],r['revision'],{},[{'asset':'a0','box':[10,20,60,45]}]);r=wait(store,j['id']);assert r['exportResult'],r
    out,_=store.output(j['id']);img=Image.open(out)
    assert not img.info and img.crop((10,20,60,45)).getbbox() is None and img.getpixel((1,1))==(255,255,255)
    assert img.getpixel((60,20))==(255,255,255) and img.getpixel((10,45))==(255,255,255)
    assert p.read_bytes()==raw
    with pytest.raises(DocumentError):store.export(j['id'],r['revision'],{},[{'asset':'a0','box':[-1,0,10,10]}])


def test_pdf_redaction_removes_data_not_overlay(store,tmp_path):
    import pymupdf as fitz
    p=tmp_path/'source.pdf';doc=fitz.open();page=doc.new_page(width=300,height=150);page.insert_text((30,60),'010-0000-1234');doc.set_metadata({'author':'김민수'});doc.save(p);doc.close()
    raw=p.read_bytes();j=store.upload(p.name,raw,p);r=wait(store,j['id']);assert r['state']=='ready',r
    assert any(c['text']=='010-0000-1234' for c in r['candidates'])
    store.export(j['id'],r['revision'],{c['id']:{'mode':'full'} for c in r['candidates']},[]);done=wait(store,j['id']);assert done['exportResult'],done
    out,_=store.output(j['id']);check=fitz.open(out)
    assert '010-0000-1234' not in ''.join(page.get_text() for page in check) and not check.metadata['author']
    assert p.read_bytes()==raw


def test_csv_roundtrip_quoted_fields_and_formula_escape(tmp_path):
    p=tmp_path/'source.csv';p.write_bytes('이름,메모\n김민수,"쉼표, 줄\n바꿈"\n=1+1,안내\n'.encode('utf-8'))
    doc=Document(p,tmp_path).load();target=next(p for p in doc.parts if p['text']=='김민수');out=tmp_path/'copy.csv';doc.export(out,{target['id']:'김**'},[])
    again=Document(out,tmp_path).load();texts=[p['text'] for p in again.parts]
    assert '김**' in texts and '쉼표, 줄\n바꿈' in texts and "'=1+1" in texts


def test_jpeg_in_word_is_losslessly_replaced_with_png(tmp_path):
    from PIL import Image
    from docx import Document as Word
    pic=tmp_path/'image.jpg';Image.new('RGB',(60,50),'white').save(pic)
    source=tmp_path/'source.docx';d=Word();d.add_picture(str(pic));d.save(source)
    doc=Document(source,tmp_path,verify=True).load();a=doc.assets[0]
    out=tmp_path/'copy.docx';doc.export(out,{},[{'asset':a['id'],'box':[5,5,25,25]}])
    loaded=Word(out);assert len(loaded.inline_shapes)==1
    entries=read_zip(out);media=[(n,b) for n,b in entries.items() if '/media/' in n]
    assert len(media)==1 and media[0][0].endswith('.png')
    img=Image.open(io.BytesIO(media[0][1]));assert img.format=='PNG'
    assert img.crop((5,5,25,25)).getbbox() is None


def test_manual_text_edit_rescanned_and_unknown_target_rejected(store):
    j=store.upload('manual.txt','미탐된 문구는 사용자가 직접 고칩니다.'.encode())
    ready=wait(store,j['id']);pid=ready['manifest']['parts'][0]['id']
    with pytest.raises(DocumentError):store.export(j['id'],ready['revision'],{},[],{'missing':'x'})
    store.export(j['id'],ready['revision'],{},[],{pid:'검토 완료'})
    done=wait(store,j['id']);assert done['exportResult'],done.get('error')
    assert store.output(j['id'])[0].read_text('utf-8')=='검토 완료'


def test_bridge_candidate_rechecks_independent_parts(store):
    base=RegexDetector()
    class Bridge:
        def check(self,title,body):
            if '\n\n' in body:return {'candidates':[{'start':0,'end':utf16(body),'source':'body','id':'bridge','label':'PERSON','text':body,'labelName':'성명','risk':'HIGH'}]}
            return base.check(title,body)
    store.detector=Bridge()
    j=store.upload('cells.csv','이름,김민수\n담당,이지은'.encode());done=wait(store,j['id'])
    assert done['state']=='ready' and {c['text'] for c in done['candidates']}=={'김민수','이지은'}


def test_cancel_queued_and_running_work_cleans_copies(store,monkeypatch):
    import threading
    started=threading.Event()
    def blocked(job,*args,**kwargs):
        started.set()
        while not job['cancel'].is_set(): time.sleep(.01)
        raise DocumentError('cancelled')
    monkeypatch.setattr(store,'_adapter',blocked)
    a=store.upload('a.txt',b'first');assert started.wait(2)
    b=store.upload('b.txt',b'second')
    work_a=store._job(a['id'])['work'];work_b=store._job(b['id'])['work']
    store.remove(b['id']);store.remove(a['id'])
    deadline=time.monotonic()+3
    while work_a.exists() and time.monotonic()<deadline:time.sleep(.02)
    assert not work_a.exists() and not work_b.exists()


def test_failed_save_does_not_publish_partial_file(store,tmp_path,monkeypatch):
    import backend.file_review as mod
    j=store.upload('sample.txt','김민수'.encode());ready=wait(store,j['id'])
    store.export(j['id'],ready['revision'],{c['id']:{'mode':'full'} for c in ready['candidates']},[])
    assert wait(store,j['id'])['exportResult']
    target=tmp_path/'out.txt'
    def fail(*args):raise OSError('simulated')
    with monkeypatch.context() as m:
        m.setattr(mod.os,'rename' if mod.os.name=='nt' else 'link',fail)
        with pytest.raises(DocumentError):store.save_new(j['id'],target)
    assert not target.exists() and not list(tmp_path.glob('.before-post-copy-*.tmp'))


def test_excel_hidden_validation_prompts_are_not_silently_ignored(tmp_path):
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    w=Workbook();w.active['A1']='안내';dv=DataValidation(type='list',formula1='"김민수,이지은"');w.active.add_data_validation(dv);dv.add('A1')
    source=tmp_path/'validation.xlsx';w.save(source)
    with pytest.raises(FormatError):Document(source,tmp_path).load()
