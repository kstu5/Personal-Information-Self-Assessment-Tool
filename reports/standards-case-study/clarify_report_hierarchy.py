from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn

root=Path(__file__).resolve().parent
src=root/'생활속에_표준적용_사례조사분석서_최종.docx'
out=root/'생활속에_표준적용_사례조사분석서_가독성수정.docx'
d=Document(src)

style=d.styles['Heading 1']
style.font.size=Pt(21);style.font.bold=True
style.paragraph_format.space_before=Pt(18);style.paragraph_format.space_after=Pt(12)
style.paragraph_format.left_indent=Cm(0)
style=d.styles['Heading 2']
style.font.size=Pt(15);style.font.bold=True
style.paragraph_format.left_indent=Cm(0)
style.paragraph_format.space_before=Pt(18);style.paragraph_format.space_after=Pt(9)
style.paragraph_format.keep_with_next=True
for n in ['Normal','List Bullet']:
 d.styles[n].font.size=Pt(11);d.styles[n].font.bold=False

# Conventional decimal subsection numbers make the heading level apparent.
for p in d.paragraphs:
 if p.style.name=='Heading 1':
  m=re.match(r'0?(\d+)\s+(.*)',p.text)
  if m:p.text=f'{int(m[1])}  {m[2]}'
 elif p.style.name=='Heading 2':
  m=re.match(r'(\d+)\s+(\d+)\s+(.*)',p.text)
  if m:p.text=f'{m[1]}.{m[2]}  {m[3]}'
 if p.style.name in ['Heading 1','Heading 2']:
  p.paragraph_format.left_indent=Cm(0)
  for r in p.runs:r.font.size=None;r.font.bold=None;r.font.color.rgb=RGBColor(0,0,0)

# Move scope information before each standard-list heading as regular prose.
for prefix,heading,text in [
 ('대상 구분','3.1','패시브 스피커의 음향 성능과 앰프 내장·무선 스피커의 음향 출력 및 전송 호환성을 구분하여 살펴본다.'),
 ('적용 범위','4.1','비접촉식 학생증 태그와 BLE 비콘 방식을 중심으로 살펴본다. 전자 출결의 인식·통신 표준과 운영 조직의 정보보호 관리표준을 구분하여 검토한다.'),
]:
 p=next(p for p in d.paragraphs if p.text.startswith(prefix))
 anchor=next(p for p in d.paragraphs if p.style.name=='Heading 2' and p.text.startswith(heading))
 p.style=d.styles['Normal'];p.clear();p.add_run(text)
 np=p._p.get_or_add_pPr().find(qn('w:numPr'))
 if np is not None:np.getparent().remove(np)
 p.paragraph_format.left_indent=Cm(0);p.paragraph_format.first_line_indent=Cm(0)
 p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(10)
 p.paragraph_format.keep_with_next=True
 anchor._p.addprevious(p._p)
anchor=next(p for p in d.paragraphs if p.style.name=='Heading 2' and p.text.startswith('2.1'))
p=d.add_paragraph('가정용 냉장고의 안전성과 성능·소비에너지를 중심으로 관련 국내외 표준을 살펴본다.')
p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(10)
anchor._p.addprevious(p._p)

# Lists use one clear typographic level: identifier, then description below it.
# Other narrative bullet labels no longer look like small headings.
is_list=False
standard_rows=0
for p in d.paragraphs:
 if p.style.name=='Heading 1':is_list=False
 if p.style.name=='Heading 2':is_list=('목록' in p.text)
 if p.style.name!='List Bullet':continue
 p.paragraph_format.left_indent=Cm(.65)
 p.paragraph_format.first_line_indent=Cm(-.3)
 p.paragraph_format.space_before=Pt(0)
 p.paragraph_format.space_after=Pt(9 if is_list else 7)
 p.paragraph_format.keep_together=True
 if is_list:
  first,sep,desc=p.text.partition('  ')
  if not sep:continue
  kind=''
  for prefix,label in [('국제 구표준 ','국제 구표준'),('국내 ','국내 표준'),('국제 ','국제표준'),('산업계 ','산업계 규격')]:
   if first.startswith(prefix):kind=label;first=first[len(prefix):];break
  if not kind:continue
  p.clear();r=p.add_run(first);r.bold=True
  p.add_run(' — '+kind)
  p.add_run().add_break()
  p.add_run(desc)
  standard_rows+=1
 else:
  for r in p.runs:r.bold=False;r.font.size=None;r.font.color.rgb=RGBColor(0,0,0)

# Keep the front-page overview label quieter than numbered subsections.
p=next(p for p in d.paragraphs if p.text=='조사 대상 한눈에 보기')
for r in p.runs:r.font.size=Pt(12);r.bold=True

d.save(out)
check=Document(out)
text='\n'.join(p.text for p in check.paragraphs)
assert '구성 안내' not in text
assert '대상 구분  ' not in text and '적용 범위  ' not in text
assert all(any(p.style.name=='Heading 2' and p.text.startswith(n) for p in check.paragraphs) for n in ['2.1','3.1','4.1'])
assert all(any(p.text.startswith(n+'   ') for p in check.paragraphs) for n in ['이름','학번','학과'])
before=set(r.target_ref for r in Document(src).part.rels.values() if 'hyperlink' in r.reltype)
after=set(r.target_ref for r in check.part.rels.values() if 'hyperlink' in r.reltype)
assert before==after
assert standard_rows==17, standard_rows
print(out)
print(f'Updated {standard_rows} standard entries; preserved all source links and personal-information blanks.')
