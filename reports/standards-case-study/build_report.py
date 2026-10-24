from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'생활속에_표준적용_사례조사분석서.docx'
d=Document()
s=d.sections[0]
s.page_width=Cm(21); s.page_height=Cm(29.7)
s.top_margin=Cm(1.7); s.bottom_margin=Cm(1.7)
s.left_margin=Cm(2); s.right_margin=Cm(2)
s.footer_distance=Cm(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet']:
 st=d.styles[name]; st.font.name='맑은 고딕'; st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'맑은 고딕')
st=d.styles['Normal']; st.font.size=Pt(10.5)
st.paragraph_format.line_spacing=1.16; st.paragraph_format.space_after=Pt(4)
for name,size in [('Title',21),('Heading 1',15),('Heading 2',11.5)]:
 st=d.styles[name];st.font.size=Pt(size);st.font.bold=True
 st.paragraph_format.space_before=Pt(10 if name!='Title' else 0)
 st.paragraph_format.space_after=Pt(5)
 st.paragraph_format.keep_with_next=True
st=d.styles['List Bullet']; st.paragraph_format.left_indent=Cm(.42);st.paragraph_format.first_line_indent=Cm(-.24)
st.paragraph_format.space_after=Pt(4)
def title(t): d.add_paragraph(t,'Title')
def h(t): d.add_paragraph(t,'Heading 1')
def sub(t): d.add_paragraph(t,'Heading 2')
def b(t):
 p=d.add_paragraph(style='List Bullet')
 if '｜' in t:
  lead,body=t.split('｜',1);p.add_run(lead+'  ').bold=True;p.add_run(body)
 else:p.add_run(t)
 p.paragraph_format.keep_together=True
 return p
def page():d.add_page_break()
def link(p,label,url):
 e=OxmlElement('w:hyperlink');e.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
 r=OxmlElement('w:r');props=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'245482');props.append(c)
 r.append(props);t=OxmlElement('w:t');t.text=label;r.append(t);e.append(r);p._p.append(e)

title('생활 속 표준 적용 사례 조사분석서')
p=d.add_paragraph('냉장고 스피커 전자 출결 서비스','Subtitle');p.paragraph_format.space_after=Pt(9)
h('1 조사 개요')
b('조사 목적｜일상생활에서 사용하는 제품과 서비스의 국내외 표준 목록 및 핵심 내용을 정리하고, 표준이 안전성·성능 비교·기기 호환성·정보보호에 기여하는 방식 분석.')
b('조사 범위｜가정용 냉장고, 패시브 및 무선 스피커, 비접촉식 학생증·BLE 비콘 기반 전자 출결 서비스. QR코드 방식은 비교 사례로 검토.')
b('분류 기준｜국가표준인 KS, ISO·IEC 국제표준, Bluetooth SIG 산업계 규격을 구분. 관련 표준의 존재와 특정 제품·서비스의 실제 채택 또는 인증 여부를 구분.')
b('핵심 관점｜측정·통신 조건의 공통화가 비교 가능성과 호환성의 기반을 제공함. 실제 성능과 보안 수준은 구현 방식 및 운영 환경까지 고려하여 평가할 필요.')
h('2 냉장고의 표준 적용 사례')
sub('2 1 국내외 표준 목록')
b('국내 KS C IEC 62552｜가정용 냉장기기－특성 및 시험 방법. 냉장기기의 필수 특성과 시험방법 규정. 최종 개정확인일 2024년 5월 28일. [1]')
b('국내 KS C IEC 60335-2-24｜가정용 및 이와 유사한 전기기기의 안전성 제2-24부. 전기 냉장고·아이스크림 기기·제빙기의 개별 안전 요구사항. IEC 제7.2판과 일치 채택. [1]')
b('국제 IEC 60335-2-24:2025｜냉장기기 등의 안전 요구사항. IEC 60335-1의 일반 안전 요구사항과 함께 사용하는 개별 제품 안전표준. [2]')
b('국제 IEC 62552 시리즈｜제1부 일반 요구사항, 제2부 성능 요구사항, 제3부 에너지 소비 및 용적. 기능별로 시험조건과 평가방법을 체계화한 표준군. [3]')
sub('2 2 주요 표준 요약 및 생활 속 의미')
b('안전성｜감전·과열 등 가정용 전기기기의 위험을 다루는 안전 평가 체계. 가연성 냉매를 사용하는 압축식 냉장기기도 안전표준의 적용 범위에 포함. [2]')
b('성능과 에너지｜냉각 성능, 소비에너지, 용적 등을 공통 방법으로 평가. IEC 62552-3:2015는 에너지 시험의 주위 온도로 16℃와 32℃를 제시하고, 사용·기후 조건에 따른 소비에너지 추정을 다룸. [3]')
b('사례 분석｜제조사가 다른 냉장고를 비교할 때 동일한 시험조건이 필요. 시험 소비전력은 비교의 기준이며, 실제 전력 사용량은 설치 환경과 문 개폐 등 사용 조건에 따라 달라질 수 있음.')
b('해석 유의점｜KS와 최신 IEC의 채택 판이 같다고 단정할 수 없음. 용적 역시 해당 판의 산정 규칙을 따라야 하며, 모든 선반·바스켓 부피를 일률적으로 빼는 방식으로 설명하는 것은 부정확. [1][3]')

page();h('3 스피커의 표준 적용 사례')
sub('3 1 조사 대상과 국내외 표준 목록')
b('대상 구분｜외부 앰프가 필요한 패시브 스피커와 앰프 내장 액티브 스피커를 구분. 블루투스 스피커는 음향 성능과 무선 전송 규격을 함께 검토할 필요.')
b('국내 KS C IEC 60268-5｜음향 시스템 장비－제5부 스피커. 수동소자로 취급되는 스피커에 적용하며, 증폭기 내장 스피커는 제외. 대응 국제표준은 IEC 60268-5 제3.1판. [4]')
b('국제 IEC 60268-5｜Sound system equipment — Part 5: Loudspeakers. 패시브 스피커의 특성 및 측정에 관련된 표준. [4]')
b('국제 IEC 60268-21:2018｜Sound system equipment — Part 21: Acoustical (output-based) measurements. 패시브·액티브 음향 시스템의 입력과 실제 음향 출력 사이의 전달 특성을 측정하는 방법 규정. [5]')
b('산업계 Bluetooth Core Specification 및 A2DP｜Bluetooth SIG가 제정한 무선 통신 및 오디오 배포 규격. A2DP는 오디오 전송에 필요한 기기 간 상호운용 기능과 절차를 규정. [6]')
sub('3 2 주요 표준 요약')
b('주파수 응답｜입력 주파수가 달라질 때 출력 음압이 어떻게 달라지는지 평가하는 특성. 재생 가능한 최저·최고 주파수 숫자만으로 음질을 판단하기보다 측정조건과 응답 변화까지 살펴볼 필요.')
b('정격 임피던스와 입력｜임피던스는 교류 신호에 대한 전기적 부하 특성으로 단순한 직류 저항과 구분. 정격 입력은 정해진 시험조건에서 해석해야 하며, 실제 사용 시 앰프의 허용 부하와 출력 조건을 함께 확인할 필요.')
b('음향 출력 측정｜측정 위치, 입력 신호 및 동작 수준에 따른 출력 특성 평가. IEC 60268-21은 작은 신호와 큰 신호 영역을 포함하여 실제 음향 출력에 기반한 물리적 측정을 다룸. [5]')
b('무선 오디오 호환성｜A2DP의 공통 코덱인 SBC를 기반으로 기본 상호운용성 확보. AAC 등 선택 코덱은 송신기와 수신기의 공통 지원 여부에 따라 사용 가능. [6]')
sub('3 3 생활 속 적용과 분석')
b('적용 사례｜스마트폰에서 블루투스 스피커로 음악을 전송할 때 공통 프로파일과 코덱 지원이 연결·재생의 기반이 됨. 소리 자체의 특성은 스피커 구조와 신호 처리 등의 영향을 받음.')
b('표준화 효과｜측정조건과 표시 의미를 공통화하여 제품 비교의 근거 제공. 무선 규격은 서로 다른 제조사의 기기 사이에서 오디오 전송이 가능하도록 지원.')
b('적용 한계｜IEC 60268-5를 앰프 내장 블루투스 스피커 전체에 그대로 적용할 수 없음. 해당 제품의 음향 출력에는 IEC 60268-21과 같은 적용 범위가 맞는 표준 검토 필요. [4][5]')
b('성능 판단｜A2DP 지원만으로 동일한 음질이나 최소 지연시간이 보장되는 것은 아님. 코덱, 버퍼 처리, 기기 성능, 전파 환경 등을 함께 고려할 필요.')

page();h('4 전자 출결 서비스의 표준 적용 사례')
sub('4 1 서비스 구성과 국내외 표준 목록')
b('서비스 구성｜학생 식별 수단, 카드·스마트폰과 단말기의 통신, 서버의 출결 판정, 기록 저장 및 조회로 구성. 통신 성공과 실제 출석의 확인은 서로 다른 단계.')
b('국내 KS X ISO/IEC 14443-2｜비접촉식 근접 카드의 RF 전력 및 신호 인터페이스 규정. 카드와 단말기 간 전력 공급·양방향 통신 특성에 관련. 2025년 확인 표준의 대응 국제표준은 ISO/IEC 14443-2:2016. [7]')
b('국내 KS X ISO/IEC 14443-3｜개인식별용 카드 및 보안장치의 초기화와 충돌방지 규정. 여러 카드가 있는 환경에서 카드 인식 절차를 다룸. ISO/IEC 14443-3:2018과 일치 채택. [7]')
b('국제 ISO/IEC 14443-2 및 14443-3｜비접촉식 근접 객체의 RF 신호 인터페이스, 초기화 및 충돌방지에 관한 표준. 해당 방식의 학생증과 단말기를 사용하는 경우 관련됨. [8]')
b('산업계 Bluetooth Core Specification의 BLE 규격｜저전력 무선 통신의 구조·절차·프로토콜 규정. 비콘 신호를 스마트폰이 수신하는 방식의 기술적 기반. [6]')
b('국제 ISO/IEC 27001:2022｜정보보호 관리체계 요구사항. 조직의 위험관리 과정을 통해 정보의 기밀성·무결성·가용성을 관리하는 표준으로, 서비스 운영 조직의 정보보호 분석에 활용. [9]')
sub('4 2 주요 표준 요약과 적용 사례')
b('학생증 태그 방식｜호환되는 카드와 단말기 사이에서 공통 통신 규칙으로 식별정보를 처리. 모든 RFID·NFC 장치가 ISO/IEC 14443 방식인 것은 아니므로 실제 카드 사양 확인 필요. [7][8]')
b('BLE 비콘 방식｜강의실의 비콘 신호를 앱이 수신하여 출결 판정에 활용. 실제 대학의 전자출결 안내에서도 비콘 및 인증번호 방식 확인 가능. [10]')
b('위치 판단 한계｜RSSI는 거리·근접성을 추정하는 수단이며, 벽·사람·간섭 등 환경의 영향을 받음. RSSI만으로 강의실 안팎의 위치를 항상 정밀하게 판별한다고 볼 수 없음. [6]')
b('정보보호 관리｜출결 기록의 무단 열람·변조·소실 위험을 평가하고 접근 권한, 기록 관리, 백업 등 통제의 필요성 검토. ISO/IEC 27001은 특정 OTP나 출석 알고리즘을 일률적으로 지정하는 표준이 아님. [9]')
b('보완 설계 예시｜유효시간이 있는 인증정보, 재사용 제한, 본인 확인, 오류 정정 신청 절차 등을 서비스 요구에 맞게 검토. 이러한 조치는 출결 서비스 설계 예시이며 개별 통신표준의 의무사항과 구분.')

page();h('5 비교 분석과 시사점')
sub('5 1 전자 출결의 다른 구현 방식')
b('QR코드 방식｜KS X ISO/IEC 18004와 ISO/IEC 18004가 QR코드의 데이터 부호화, 기호 형식, 오류 정정, 판독 등에 관한 공통 규칙 제공. 실제로 한국외대는 학교 앱으로 QR코드를 촬영하는 출석 확인 절차를 안내. [10][11]')
b('국내외 판의 차이｜KS X ISO/IEC 18004의 2022년 확인판은 ISO/IEC 18004:2000과 일치 채택으로 표시되며, ISO의 2024년판과 구분 필요. [11]')
b('RFID 방식의 구분｜ISO/IEC 18000-63은 860~960 MHz 대역의 UHF RFID 무선 인터페이스를 다룸. 이를 근접 학생증이나 BLE 출결에 공통 적용되는 표준으로 취급하는 것은 부적절. [12]')
sub('5 2 세 대상의 표준화 목적 비교')
b('냉장고｜안전 위험의 관리와 성능·소비에너지의 객관적 비교에 중점. 시험조건이 달라지면 소비전력이나 용적 수치의 직접 비교에 한계 발생.')
b('스피커｜음향 특성 측정과 오디오 전송 호환성에 중점. 소리의 물리적 특성과 기기 간 통신 가능 여부를 구분하여 분석할 필요.')
b('전자 출결｜식별·통신의 호환성과 기록의 신뢰성에 중점. 표준 통신을 사용하더라도 대리 출석, 계정 공유, 통신 장애 등의 문제에는 별도 대응 필요.')
b('공통 효과｜제조사·운영기관 간 공통 기술 언어 제공, 시험·평가의 재현성 향상, 이용자의 정보 비교 및 서비스 이용 편의 지원.')
sub('5 3 표준 적용 시 유의사항')
b('범위와 판 확인｜제품 유형, 통신 방식, 표준의 발행연도 및 개정 상태 확인 필요. 국제표준과 이를 채택한 KS 사이에는 판의 차이가 존재할 수 있음.')
b('표준과 인증 구분｜표준은 요구사항이나 평가방법을 정한 문서이며, 인증은 정해진 범위에서 적합성을 확인하는 절차. 관련 표준을 열거하는 것만으로 실제 제품의 인증 사실이 입증되지는 않음.')
b('적용 증거 확인｜특정 제품·서비스의 실제 준수 여부는 제품 설명서, 시험성적서, 인증서 또는 운영기관 기술자료로 확인할 필요. 본인 확인과 강의 출석의 사실 확인도 구분 필요.')
h('6 결론')
b('냉장고와 스피커의 사례를 통해 표준화가 안전 평가와 성능 비교의 공통 기준을 마련한다는 점 확인.')
b('전자 출결의 사례를 통해 여러 통신표준과 정보보호 관리체계가 결합하여 서비스의 기술적 기반을 구성한다는 점 확인.')
b('생활 속 표준의 가치는 공통 요구사항과 평가방법의 제공에 있음. 효과적인 적용을 위해서는 표준의 범위에 맞는 구현, 사용 환경을 반영한 평가, 지속적인 운영 관리가 함께 필요.')

page();h('7 참고문헌')
p=d.add_paragraph('공식 표준기관 및 서비스 운영기관 자료  ·  자료 확인일 2026년 10월 21일')
p.runs[0].font.size=Pt(9)
refs=[
('[1] 국가기술표준원 e나라 표준인증. 냉장기기 국내 표준.',[
 ('KS C IEC 62552','https://standard.go.kr/KSCI/standardIntro/getStandardSearchView.do?ksNo=KSCIEC62552&menuId=503&reformNo=02&tmprKsNo=KSCIEC62552&topMenuId=502'),
 ('KS C IEC 60335-2-24','https://www.standard.go.kr/KSCI/standardIntro/getStandardSearchView.do?ksNo=KSCIEC60335-2-24&menuId=919&reformNo=&tmprKsNo=KSCIEC60335-2-24&topMenuId=502&upperMenuId=503')]),
('[2] IEC. IEC 60335-2-24:2025. 냉장기기 등의 안전 요구사항.',[('표준 상세','https://webstore.iec.ch/en/publication/75938')]),
('[3] IEC. IEC 62552 시리즈. 냉장기기 특성 및 시험방법.',[
 ('제2부 성능 및 시리즈 구성','https://webstore.iec.ch/en/publication/21804'),('제3부 에너지 및 용적','https://webstore.iec.ch/en/publication/21803'),('제3부 개정 이력','https://webstore.iec.ch/en/publication/61968')]),
('[4] 국가기술표준원 및 IEC. 스피커 표준.',[
 ('KS C IEC 60268-5','https://standard.go.kr/KSCI/standardIntro/getStandardSearchView.do?ksNo=KSCIEC60268-5&menuId=919&pageIndex=1&pageUnit=10&reformNo=19&tmprKsNo=KSCIEC60268-5&topMenuId=502&upperMenuId=503'),('IEC 60268-5','https://webstore.iec.ch/en/publication/1223')]),
('[5] IEC. IEC 60268-21:2018. 음향 출력 기반 측정.',[('표준 상세','https://webstore.iec.ch/en/publication/28687')]),
('[6] Bluetooth SIG. 무선 통신 및 오디오 배포 규격과 기술 해설.',[
 ('규격 목록','https://www.bluetooth.com/specifications/specs/'),('A2DP','https://www.bluetooth.com/specifications/specs/advanced-audio-distribution-profile/'),
 ('BLE 입문 자료','https://www.bluetooth.com/bluetooth-le-primer/'),('RSSI의 한계','https://www.bluetooth.com/core-specification-6-feature-overview/'),
 ('A2DP 코덱 적합성 항목','https://files.bluetooth.com/download/a2dp-ics-p16ed3-pdf/')]),
('[7] 한국표준협회. 비접촉식 근접 카드 국내 표준 및 부합화 정보.',[
 ('KS X ISO/IEC 14443-2','https://kssn.net/search/stddetail.do?itemNo=K001010128795'),('KS X ISO/IEC 14443-3','https://www.kssn.net/search/stddetail.do?itemNo=K001010156652')]),
('[8] ISO. 비접촉식 근접 객체 국제표준.',[
 ('ISO/IEC 14443-2:2020','https://www.iso.org/standard/73597.html'),('ISO/IEC 14443-3:2018','https://www.iso.org/standard/73598.html?browse=tc')]),
('[9] ISO. ISO/IEC 27001:2022. 정보보호 관리체계 요구사항.',[('표준 및 관리체계 설명','https://committee.iso.org/standard/27001')]),
('[10] 대학 전자출결 운영 안내.',[
 ('아주대학교 비콘 및 인증번호 방식','https://security.ajou.ac.kr/kr/ajou/notice.do?article.offset=0&articleLimit=10&articleNo=373743&mode=view'),('한국외국어대학교 QR코드 방식','https://at.hufs.ac.kr/guide.html?display=eng')]),
('[11] 한국표준협회 및 ISO. QR코드 기호 사양.',[
 ('KS X ISO/IEC 18004 2022 확인','https://www.kssn.net/search/stddetail.do?itemNo=K001010137501'),('ISO/IEC 18004:2024','https://www.iso.org/standard/83389.html')]),
('[12] ISO. ISO/IEC 18000-63:2021. 860~960 MHz Type C RFID.',[('표준 상세','https://www.iso.org/standard/78309.html')]),
]
for label,ls in refs:
 p=d.add_paragraph();p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(2);p.paragraph_format.keep_with_next=True
 r=p.add_run(label);r.font.size=Pt(9);r.bold=True
 p=d.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.keep_together=True;p.paragraph_format.line_spacing=1.1
 for j,(lab,url) in enumerate(ls):
  if j:p.add_run('  |  ').font.size=Pt(9)
  link(p,lab,url)
 for r in p._p.iter(qn('w:r')):
  pr=r.find(qn('w:rPr'))
  if pr is None:pr=OxmlElement('w:rPr');r.insert(0,pr)
  sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');pr.append(sz)

f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=f.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
d.core_properties.title='생활 속 표준 적용 사례 조사분석서'
d.core_properties.subject='냉장고 스피커 전자 출결 서비스의 국내외 표준 및 적용 분석'
d.core_properties.author='';d.core_properties.keywords='표준화, KS, IEC, ISO, Bluetooth, 전자 출결'
d.save(OUT)
print(OUT)
