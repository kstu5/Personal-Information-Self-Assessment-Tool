from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

root=Path(__file__).resolve().parent
source=root/'생활속에_표준적용_사례조사분석서_디자인_수정.docx'
out=root/'생활속에_표준적용_사례조사분석서_최종.docx'
d=Document(source)
bullet_properties=deepcopy(next(p for p in d.paragraphs if p.style.name=='List Bullet')._p.pPr)
def insert(anchor,text,style='List Bullet'):
 p=d.add_paragraph(style=style)
 if style=='List Bullet':
  old=p._p.pPr
  if old is not None:p._p.remove(old)
  p._p.insert(0,deepcopy(bullet_properties))
  lead,sep,body=text.partition('｜')
  if sep:
   r=p.add_run(lead+'  ');r.bold=True;r.font.color.rgb=RGBColor.from_string('202938');p.add_run(body)
  else:p.add_run(text)
 else:p.add_run(text)
 anchor._p.addprevious(p._p)
 return p
def section(number):return next(p for p in d.paragraphs if p.style.name=='Heading 1' and p.text.startswith(number))
def replace_between(start,finish,items):
 a=section(start);z=section(finish);node=a._p.getnext()
 while node is not z._p:
  nxt=node.getnext();node.getparent().remove(node);node=nxt
 for style,text in items:insert(z,text,style)
def B(text):return ('List Bullet',text)
def H(text):return ('Heading 2',text)

# Replace the four outline-like overview bullets with the agreed two paragraphs.
overview=[
'일상생활에서 사용하는 제품과 서비스에는 안전성, 성능, 기기 간 호환성 등을 확보하기 위한 다양한 표준이 적용된다. 표준은 제품의 특성을 일정한 조건에서 측정하고 비교할 수 있는 기준을 제공하며, 서로 다른 기기와 시스템이 정보를 주고받는 데 필요한 공통 규칙으로도 활용된다. 따라서 생활 속 제품과 서비스의 관련 표준을 살펴보는 것은 우리가 이용하는 기술의 작동 방식과 신뢰성의 기반을 이해하는 데 의미가 있다.',
'본 보고서는 냉장고, 스피커, 전자 출결 서비스를 대상으로 국내외 관련 표준의 목록과 주요 내용을 조사하고, 각 표준이 실제 이용 과정에서 어떤 역할을 하는지 분석한다. 냉장고는 안전성과 에너지 소비, 스피커는 음향 성능과 무선 연결의 호환성, 전자 출결 서비스는 사용자 식별과 통신 및 정보보호를 중심으로 살펴본다. 이를 바탕으로 제품과 서비스의 특성에 따른 표준화 목적의 차이와 적용 시 고려해야 할 한계를 비교하고자 한다.'
]
anchor=next(p for p in d.paragraphs if p.text=='조사 대상 한눈에 보기')
start=section('01');node=start._p.getnext()
while node is not anchor._p:
 nxt=node.getnext();node.getparent().remove(node);node=nxt
for text in overview:
 p=insert(anchor,text,'Normal');p.paragraph_format.space_after=Pt(10)
 p.paragraph_format.first_line_indent=Pt(0)

cold=[
H('2 1   국내외 표준 목록'),
B('국내 KS C IEC 62552｜가정용 냉장기기 특성 및 시험 방법. 최종 개정확인일 2024년 5월 28일. [1]'),
B('국내 KS C IEC 60335-2-24｜전기 냉장고·아이스크림 기기·제빙기의 개별 안전 요구사항. 대응 국제표준 IEC 제7.2판과 일치 채택. [1]'),
B('국제 IEC 60335-2-24:2025｜가정용 및 이와 유사한 전기기기의 안전성 제2-24부. 냉장기기 등의 개별 요구사항. [2]'),
B('국제 IEC 62552-1:2015｜가정용 냉장기기 특성 및 시험방법 제1부 일반 요구사항. [3]'),
B('국제 IEC 62552-2:2015｜가정용 냉장기기 특성 및 시험방법 제2부 성능 요구사항. [3]'),
B('국제 IEC 62552-3:2015｜가정용 냉장기기 특성 및 시험방법 제3부 에너지 소비 및 용적. 제1~3부에는 2020년 및 2025년 개정이 있으며, 적용 시 해당 개정 포함 여부 확인 필요. [3]'),
H('2 2   표준별 요약문'),
B('KS C IEC 62552｜내부 자연대류 또는 강제 공기순환 방식으로 냉각되는 가정용 냉장기기의 필수 특성과 이를 확인하는 시험방법을 규정. 제조사 간 성능을 공통 조건에서 평가하기 위한 국내 기준으로 활용. [1]'),
B('KS C IEC 60335-2-24｜전기 냉장고 등 대상 기기의 안전을 다루는 국내 표준. 일반 안전 요구사항과 함께 해당 제품의 위험을 평가하며, 채택한 국제표준 판은 2025년 국제판과 구분 필요. [1]'),
B('IEC 60335-2-24:2025｜가정용 냉장기기 등의 안전 요구사항을 규정하며 가연성 냉매를 사용하는 압축식 기기도 포함. IEC 60335-1과 함께 적용하여 대상 기기의 공통 위험과 개별 위험을 평가. [2]'),
B('IEC 62552-1｜냉장기기의 기본 특성과 시험을 위한 일반 요구사항을 규정. 기본 설계·동작의 형식시험을 다루며, 생산품 표본추출이나 인증 절차 자체를 정하는 표준은 아님. [3]'),
B('IEC 62552-2｜냉장기기의 성능 특성을 확인하는 시험방법을 규정. 냉각 능력과 냉각시간 등 개별 성능을 공통 방법으로 평가하여 제품 비교의 근거 제공. [3]'),
B('IEC 62552-3｜소비에너지와 용적의 측정방법 및 사용·기후 조건에 따른 소비에너지 추정 방법을 규정. 2015년 기본판은 에너지 시험의 주위 온도로 16℃와 32℃를 제시. [3]'),
H('2 3   생활 속 적용과 분석'),
B('비교 가능성｜제조사가 다른 냉장고의 소비에너지·성능을 비교하려면 시험조건의 공통화가 필요. 실제 전력 사용량은 설치 환경과 문 개폐 등 사용 조건에 따라 달라질 수 있음.'),
B('해석 유의점｜KS와 최신 IEC의 채택 판이 같다고 단정할 수 없음. 용적은 해당 판의 산정 규칙을 따라야 하며, 모든 선반·바스켓 부피를 일률적으로 빼는 방식으로 설명하는 것은 부정확. [1][3]'),
]
replace_between('02','03',cold)

speaker=[
H('3 1   국내외 표준 및 규격 목록'),
B('대상 구분｜패시브 스피커의 음향 성능과 앰프 내장·무선 스피커의 음향 출력 및 전송 호환성을 구분하여 검토.'),
B('국내 KS C IEC 60268-5｜음향 시스템 장비 제5부 스피커. 패시브 스피커 대상이며 증폭기 내장 스피커는 제외. IEC 제3.1판과 일치 채택. [4]'),
B('국제 구표준 IEC 60268-5:2003+A1:2007｜Sound system equipment — Part 5: Loudspeakers. 국내 KS의 대응 국제표준이며, IEC에서는 2026년 4월 17일 철회 상태로 표시. [4]'),
B('국제 IEC 60268-21:2018｜Sound system equipment — Part 21: Acoustical (output-based) measurements. 음향 출력 기반 측정 표준. [5]'),
B('산업계 Bluetooth Core Specification｜Bluetooth SIG의 무선 통신 기반 규격. ISO·IEC 국제표준과 구분되는 산업계 규격. [6]'),
B('산업계 Bluetooth A2DP｜Advanced Audio Distribution Profile. 고품질 오디오 배포를 위한 기기 기능과 절차에 관한 프로파일. [6]'),
H('3 2   표준별 요약문'),
B('KS C IEC 60268-5｜패시브 스피커의 특성과 관련 측정을 다루는 국내 표준. 서로 다른 스피커를 비교할 수 있는 기준 제공. 대응 IEC 판의 철회가 국내 KS의 자동 폐지를 뜻하는 것은 아니므로 국내 상태를 별도 확인할 필요. [4]'),
B('IEC 60268-5 구표준｜패시브 스피커에 표시할 특성과 정현파·규정 잡음·충격성 신호 등을 이용하는 측정방법을 규정했던 표준. 국내 표준의 기반을 이해하는 자료로 검토하며 현행 국제표준으로 소개하지 않음. [4]'),
B('IEC 60268-21｜패시브·액티브 음향 시스템에서 입력 신호와 실제 음향 출력 사이의 전달 특성을 물리적으로 측정하는 방법 규정. 작은 신호와 큰 신호 영역을 포함하며, 청취자가 느끼는 음질의 인지적 평가는 범위에 포함하지 않음. [5]'),
B('Bluetooth Core Specification｜기기 간 무선 통신의 구조·계층·절차·프로토콜을 정의. 스피커와 스마트폰이 통신하는 기반을 제공하며, 오디오 서비스의 세부 상호운용성은 관련 프로파일과 함께 검토. [6]'),
B('Bluetooth A2DP｜오디오 송신기와 수신기 사이의 기능·절차 및 코덱 호환성을 규정. SBC를 기본 공통 코덱으로 사용하며, AAC 등 선택 코덱은 양쪽의 지원 여부에 따라 사용 가능. [6]'),
H('3 3   생활 속 적용과 분석'),
B('음향 특성｜주파수 응답은 입력 주파수에 따른 출력 음압의 변화를 뜻함. 임피던스는 교류 신호에 대한 부하 특성이며, 정격 입력은 정해진 시험조건에서 해석할 필요.'),
B('적용 사례｜스마트폰에서 블루투스 스피커로 음악을 전송할 때 공통 프로파일과 코덱 지원이 연결·재생의 기반이 됨. 소리 자체의 특성은 스피커 구조와 신호 처리 등의 영향도 받음.'),
B('적용 한계｜IEC 60268-5의 범위를 앰프 내장 블루투스 스피커 전체로 확대할 수 없음. A2DP 지원만으로 동일한 음질이나 최소 지연시간이 보장되는 것도 아니므로 제품 사양과 사용 환경을 함께 평가.'),
]
replace_between('03','04',speaker)

attendance=[
H('4 1   국내외 표준 및 규격 목록'),
B('적용 범위｜비접촉식 학생증 태그와 BLE 비콘 방식을 중심으로 검토. 전자 출결의 인식·통신 표준과 운영 조직의 정보보호 관리표준을 구분.'),
B('국내 KS X ISO/IEC 14443-2｜비접촉식 근접 카드의 RF 전력 및 신호 인터페이스. 2025년 확인 정보의 대응 국제표준은 ISO/IEC 14443-2:2016. [7]'),
B('국내 KS X ISO/IEC 14443-3｜개인식별용 카드 및 보안장치의 초기화와 충돌방지. ISO/IEC 14443-3:2018과 일치 채택. [7]'),
B('국제 ISO/IEC 14443-2:2020｜개인 식별용 카드 및 보안 장치의 비접촉식 근접 객체 제2부 RF 전력 및 신호 인터페이스. [8]'),
B('국제 ISO/IEC 14443-3:2018｜개인 식별용 카드 및 보안 장치의 비접촉식 근접 객체 제3부 초기화와 충돌방지. [8]'),
B('산업계 Bluetooth Core Specification의 BLE 규격｜Bluetooth SIG의 저전력 무선 통신 규격. 비콘 방식의 기술적 기반. [6]'),
B('국제 ISO/IEC 27001:2022｜정보보호 관리체계 요구사항. 특정 출결 앱의 통신 규격이 아니라 운영 조직의 정보보호 관리체계에 관한 표준. [9]'),
H('4 2   표준별 요약문'),
B('KS X ISO/IEC 14443-2｜단말기와 근접 카드 사이의 전력 공급 및 양방향 통신에 필요한 RF 인터페이스 특성을 규정. 해당 방식의 학생증과 판독기 간 호환성을 검토하는 국내 기준. [7]'),
B('KS X ISO/IEC 14443-3｜카드의 인식, 통신 초기화, 여러 카드 중 하나를 선택하는 충돌방지 절차를 규정. 학생증 판독을 시작하는 단계의 기술적 기반. [7]'),
B('ISO/IEC 14443-2:2020｜근접 단말기와 카드·객체 사이에서 전력과 양방향 통신을 제공하는 장의 특성을 규정. 국가별 전자파·인체 노출 규정의 준수 방법 자체를 정하는 표준은 아님. [8]'),
B('ISO/IEC 14443-3:2018｜카드 탐색, 초기 명령과 응답, 초기 통신의 형식·타이밍 및 충돌방지 절차를 규정. 여러 카드가 있을 때 하나를 선택하여 통신을 시작하도록 지원. [8]'),
B('Bluetooth BLE 규격｜저전력 기기의 무선 신호 교환 절차와 프로토콜을 정의. 출결 앱은 비콘 신호를 판정 자료로 활용할 수 있으나, 출석 인정 조건은 서비스가 별도로 설계할 필요. [6]'),
B('ISO/IEC 27001:2022｜위험관리에 기반하여 정보의 기밀성·무결성·가용성을 관리하는 체계의 요구사항을 규정. 출결 기록의 열람 권한, 변조·소실 위험, 백업 등을 관리하는 틀로 활용 가능하며 특정 OTP 알고리즘을 일률적으로 지정하지 않음. [9]'),
H('4 3   생활 속 적용과 분석'),
B('서비스 흐름｜학생 식별, 카드·스마트폰과 단말기의 통신, 서버의 출결 판정, 기록 저장 및 조회로 구성. 실제 대학 안내에서도 비콘·인증번호 방식 확인 가능. [10]'),
B('기술적 한계｜모든 RFID·NFC 장치가 ISO/IEC 14443 방식인 것은 아니므로 실제 카드 사양 확인 필요. BLE의 RSSI 역시 환경의 영향을 받는 근접성 추정 자료이므로 강의실 내외 위치를 항상 정밀하게 판별한다고 볼 수 없음. [6][7][8]'),
B('운영 보완｜유효시간이 있는 인증정보, 재사용 제한, 본인 확인, 오류 정정 신청 등을 서비스 요구에 맞게 검토. 통신 성공과 실제 학생의 출석 확인을 구분하고, 통신표준 자체의 의무사항과 운영 설계 예시를 구별.'),
]
replace_between('04','05',attendance)

# Add an explicit paired summary for QR standards already discussed in chapter 5.
qr=next(p for p in d.paragraphs if p.text.startswith('QR코드 방식'))
qr.clear();r=qr.add_run('QR코드 방식  ');r.bold=True
qr.add_run('국내 KS X ISO/IEC 18004와 국제 ISO/IEC 18004:2024가 관련 표준. 실제로 한국외대는 학교 앱으로 QR코드를 촬영하는 출석 확인 절차를 안내. [10][11]')
qrnext=next(p for p in d.paragraphs if p.text.startswith('국내외 판의 차이'))
insert(qrnext,'KS X ISO/IEC 18004 요약｜QR코드의 데이터 부호화, 기호 형식·치수, 오류 정정 및 판독 등의 요구사항을 규정. 2022년 확인판은 ISO/IEC 18004:2000과 일치 채택으로 표시. [11]')
insert(qrnext,'ISO/IEC 18004:2024 요약｜QR코드의 특성, 문자 인코딩, 심벌 형식, 오류 정정, 참조 판독 알고리즘 및 품질 요구사항을 규정. 출결정보의 표현·인식에 관련되며, QR 이미지의 공유나 대리 출석을 표준 자체가 방지하는 것은 아님. [11]')

def append_link(p,label,url):
 p.add_run('  |  ')
 e=OxmlElement('w:hyperlink');e.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'245482');pr.append(c)
 sz=OxmlElement('w:sz');sz.set(qn('w:val'),'19');pr.append(sz);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);e.append(r);p._p.append(e)
p=next(p for p in d.paragraphs if '제2부 성능 및 시리즈 구성' in p.text)
append_link(p,'제1부 일반 요구사항','https://webstore.iec.ch/en/publication/21805')
p=next(p for p in d.paragraphs if p.text.startswith('KS C IEC 60268-5  |'))
append_link(p,'IEC 제3.1판 및 철회 정보','https://webstore.iec.ch/en/publication/1224')

# In the overview matrix, clearly distinguish the withdrawn international part.
cell=d.tables[0].rows[2].cells[2]
cell.text='IEC 60268-21\nKS C IEC 60268-5\nBluetooth A2DP'
for p in cell.paragraphs:
 p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(2)
 for r in p.runs:r.font.size=Pt(9.5)

# Keep existing style hierarchy, identity blanks and references; use natural
# pagination when added summaries need more space.
d.core_properties.comments=''
d.save(out)
check=Document(out)
alltext='\n'.join(p.text for p in check.paragraphs)
assert '구성 안내' not in alltext
assert sum(p.style.name=='Normal' and p.text in overview for p in check.paragraphs)==2
assert sum(p.style.name=='Heading 2' and '표준별 요약문' in p.text for p in check.paragraphs)==3
assert all(any(p.text.startswith(label+'   ') for p in check.paragraphs) for label in ['이름','학번','학과'])
assert all(code in alltext for code in ['IEC 62552-1','IEC 62552-2','IEC 62552-3','IEC 60268-21','ISO/IEC 14443-2','ISO/IEC 14443-3','ISO/IEC 27001','ISO/IEC 18004'])
before_links=set(r.target_ref for r in Document(source).part.rels.values() if 'hyperlink' in r.reltype)
after_links=set(r.target_ref for r in check.part.rels.values() if 'hyperlink' in r.reltype)
assert before_links<=after_links
print(out)
print('Verified prose overview, three explicit standard-summary sections, identity fields and all reference links.')
