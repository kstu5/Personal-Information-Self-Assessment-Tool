"""Bounded, location-aware document adapters. Used only inside the file worker.

Every target is derived from our private input copy, never from a client path.
Editing rewrites a new file and checks the edited text/opaque pixels before return.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import time
import zipfile

from lxml import etree as E
from PIL import Image, ImageDraw, ImageOps

ROOT = Path(__file__).resolve().parents[1]
TOOLS = ROOT/'.runtime'/'file-tools'
MAX_BYTES = 20 * 1024 * 1024
MAX_TEXT = 180_000
MAX_PARTS = 4000
MAX_PIXELS = 20_000_000
SUPPORTED = {'.txt', '.csv', '.tsv', '.docx', '.pptx', '.hwpx', '.xlsx', '.pdf', '.hwp', '.png', '.jpg', '.jpeg'}
Image.MAX_IMAGE_PIXELS = MAX_PIXELS


class FormatError(Exception):
    pass


def local_name(node):
    return E.QName(node).localname if isinstance(node.tag, str) else ''


def parse_xml(raw):
    if b'<!DOCTYPE' in raw.upper() or b'<!ENTITY' in raw.upper():
        raise FormatError('외부 엔터티가 포함된 문서는 처리할 수 없습니다.')
    return E.fromstring(raw, E.XMLParser(resolve_entities=False, no_network=True, huge_tree=False))


def read_zip(path):
    with zipfile.ZipFile(path) as z:
        infos = z.infolist()
        names = [i.filename for i in infos]
        if len(infos) > 6000 or len(set(names)) != len(names) or sum(i.file_size for i in infos) > 120*1024*1024:
            raise FormatError('문서 내부 크기 또는 항목 수가 처리 한도를 넘습니다.')
        if any(i.flag_bits & 1 or i.file_size > 32*1024*1024 or i.file_size > max(1, i.compress_size)*1000 for i in infos):
            raise FormatError('암호화되었거나 압축 크기가 비정상적인 문서입니다.')
        if any(n.startswith(('/', '\\')) or '..' in n.replace('\\', '/').split('/') for n in names):
            raise FormatError('문서 내부 경로가 올바르지 않습니다.')
        return {i.filename: z.read(i) for i in infos}


def write_zip(path, entries):
    with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as z:
        for name, raw in entries.items():
            z.writestr(name, raw, compress_type=zipfile.ZIP_STORED if name == 'mimetype' else zipfile.ZIP_DEFLATED)


class Document:
    def __init__(self, path, workspace, *, verify=False):
        self.path, self.work = Path(path), Path(workspace)
        self.kind = self.path.suffix.lower()
        self.verify = verify
        self.parts, self.assets, self.warnings, self.blockers, self.coverage = [], [], [], [], []
        self.entries, self.trees, self.workbook = {}, {}, None
        self.calculation = None

    def add(self, text, location, target, *, boxes=None, asset=None):
        text = str(text)
        if not text.strip(): return
        if len(text) > 30000 or len(self.parts) >= MAX_PARTS or sum(len(p['text']) for p in self.parts)+len(text) > MAX_TEXT:
            raise FormatError('검사 분량이 한도를 넘었습니다. 문서를 나누어 검사해 주세요. (18만 자·4천 영역)')
        p = {'id': f'p{len(self.parts)}', 'text': text, 'location': location, 'target': target}
        if boxes is not None: p['boxes'] = boxes
        if asset is not None: p['asset'] = asset
        self.parts.append(p)

    def load(self):
        if self.kind not in SUPPORTED: raise FormatError('지원하지 않는 파일 형식입니다.')
        if self.path.stat().st_size > MAX_BYTES: raise FormatError('파일당 20MB까지 검사할 수 있습니다.')
        if self.kind in {'.docx', '.pptx', '.hwpx'}: self.read_package()
        elif self.kind == '.xlsx': self.read_excel()
        elif self.kind == '.pdf': self.read_pdf()
        elif self.kind in {'.png', '.jpg', '.jpeg'}:
            self.image(self.path.read_bytes(), '이미지', {'kind': 'image'}, use_ocr=not self.verify)
            self.coverage = ['이미지 OCR(한국어·영어)', '직접 지정한 가림 영역', '새 PNG의 메타데이터 제거']
        elif self.kind == '.hwp': self.read_hwp()
        else: self.read_text()
        return self

    def manifest(self):
        return {'parts': self.parts, 'assets': self.assets, 'warnings': list(dict.fromkeys(self.warnings)),
                'blockers': list(dict.fromkeys(self.blockers)), 'coverage': self.coverage,
                'calculation': self.calculation, 'kind': self.kind,
                'outputSuffix': '.png' if self.kind in {'.png', '.jpg', '.jpeg'} else self.kind}

    def read_text(self):
        raw = self.path.read_bytes()
        for encoding in ('utf-8-sig', 'utf-16' if raw[:2] in (b'\xff\xfe', b'\xfe\xff') else 'cp949'):
            try: self.text = raw.decode(encoding); break
            except UnicodeError: continue
        else: raise FormatError('UTF-8·UTF-16·CP949로 읽을 수 없는 텍스트입니다.')
        if '\0' in self.text: raise FormatError('텍스트가 아닌 데이터가 포함되어 있습니다.')
        if self.kind in {'.csv', '.tsv'}:
            self.delimiter = '\t' if self.kind == '.tsv' else ','
            self.rows = list(csv.reader(io.StringIO(self.text, newline=''), delimiter=self.delimiter))
            for r, row in enumerate(self.rows):
                for c, value in enumerate(row): self.add(value, f'{r+1}행 · {c+1}열', {'kind': 'csv', 'row': r, 'col': c})
            self.warnings.append('CSV/TSV는 쉼표·탭 구분으로 읽고 UTF-8로 저장합니다. =, +, -, @로 시작하는 셀은 배포본에서 수식 실행 방지용 작은따옴표를 붙입니다.')
        else:
            # Split on line boundaries; never silently truncate.
            self.lines = self.text.splitlines(keepends=True)
            for i, line in enumerate(self.lines): self.add(line, f'{i+1}행', {'kind': 'line', 'index': i})
        self.coverage = ['전체 텍스트', '행·열 위치', 'UTF-8 사본']

    def read_package(self):
        self.entries = read_zip(self.path)
        required = {'.docx': 'word/document.xml', '.pptx': 'ppt/presentation.xml', '.hwpx': 'Contents/content.hpf'}[self.kind]
        if required not in self.entries: raise FormatError('확장자와 실제 문서 구조가 일치하지 않습니다.')
        suspicious = [n for n in self.entries if any(s in n.lower() for s in
            ('/embeddings/', 'vbaproject', '_xmlsignatures', '/activex/', 'threadedcomment', '/diagrams/', '/charts/'))]
        if suspicious: self.blockers.append('내장 문서·차트·매크로·서명·현대식 댓글 등 지원하지 않는 개체가 있습니다. 이 문서는 검사만 가능합니다.')
        if self.kind == '.hwpx':
            self.warnings.append('HWPX 사본은 내부 미리보기 텍스트·이미지를 제거합니다. 일부 한글 버전에서는 미리보기가 비어 보일 수 있습니다.')
        for name, raw in self.entries.items():
            if name.startswith('Preview/'): continue
            if name.lower().startswith('docprops/thumbnail'): continue
            if name.endswith(('.xml', '.rels', '.hpf')):
                tree = parse_xml(raw); self.trees[name] = tree
                if name.startswith('customXml/') and not name.endswith('.rels'):
                    # Word's empty bibliography is boilerplate, but arbitrary custom payloads are not covered.
                    safe = local_name(tree)=='Sources' and len(tree)==0 or local_name(tree)=='datastoreItem' and all(local_name(e) in {'datastoreItem','schemaRefs','schemaRef'} for e in tree.iter())
                    if not safe:self.blockers.append('사용자 지정 XML 데이터가 있어 사본 저장을 지원하지 않습니다.')
                node_ids = {e: i for i, e in enumerate(tree.iter())}
                label = self.package_location(name)
                selected = set()
                for para in tree.iter():
                    if local_name(para) != 'p': continue
                    nodes = [e for e in para.iter() if local_name(e) in {'t', 'delText', 'instrText'} and e.text and e not in selected]
                    if not nodes: continue
                    # Text boxes may contain nested paragraphs: only the nearest paragraph owns a text run.
                    nodes = [e for e in nodes if next((a for a in e.iterancestors() if local_name(a) == 'p'), None) is para]
                    if not nodes: continue
                    selected.update(nodes)
                    text = ''.join(e.text for e in nodes)
                    deleted = any(local_name(a) == 'del' for e in nodes for a in e.iterancestors())
                    memo = any(local_name(a) in {'memo','memogroup'} for a in para.iterancestors()) or any(local_name(e)=='fieldBegin' and e.get('type')=='MEMO' for e in para.iter())
                    self.add(text, f'{label} · 문단 {node_ids[para]+1}'+(' · 삭제된 변경 기록' if deleted else ' · 메모' if memo else ''),
                             {'kind': 'xml', 'name': name, 'nodes': [node_ids[e] for e in nodes]})
                for e, idx in node_ids.items():
                    ln = local_name(e)
                    string_property = name.startswith('docProps/') and ln in {'title','subject','creator','keywords','description','lastModifiedBy','category','contentStatus','identifier','Company','Manager','HyperlinkBase','lpstr','lpwstr'}
                    if e not in selected and e.text and (ln in {'t', 'delText', 'instrText', 'text','meta','metaTag'} or ln=='stringParam' and e.get('name')=='Author' or string_property):
                        self.add(e.text, f'{label} · {ln}', {'kind': 'xml', 'name': name, 'nodes': [idx]})
                    for attr, value in e.attrib.items():
                        an = E.QName(attr).localname
                        if an.lower() in {'author', 'creator', 'lastmodifiedby', 'initials', 'username'} or ln in {'cmAuthor', 'author','trackChangeAuthor'} and an == 'name':
                            self.add(value, f'{label} · 작성자', {'kind': 'attr', 'name': name, 'node': idx, 'attr': attr})
                    if ln == 'Relationship' and e.get('TargetMode') == 'External':
                        self.add(e.get('Target', ''), f'{label} · 외부 링크 주소', {'kind': 'attr', 'name': name, 'node': idx, 'attr': 'Target'})
                        if not e.get('Type', '').endswith('/hyperlink'):
                            self.blockers.append('외부에서 불러오는 그림·문서 연결이 있어 사본 저장을 지원하지 않습니다.')
                if any(local_name(e).lower() in {'altchunk', 'oleobject'} for e in tree.iter()):
                    self.blockers.append('추가 문서 또는 OLE 개체를 포함해 사본 저장을 지원하지 않습니다.')
            if '/media/' in name or name.startswith('BinData/'):
                if Path(name).suffix.lower() in {'.png', '.jpg', '.jpeg', '.bmp'}:
                    self.image(raw, f'내장 이미지 · {Path(name).name}', {'kind': 'embedded', 'name': name}, use_ocr=not self.verify)
                else:
                    self.blockers.append('지원하지 않는 내장 이미지·이진 개체가 있습니다. PNG·JPEG·BMP만 직접 가릴 수 있습니다.')
        self.coverage = ['본문·표·텍스트 상자', '머리말·꼬리말·각주', '저장된 메모·발표자 노트', '읽을 수 있는 변경 기록·작성자·문서 속성', '내장 PNG·JPEG·BMP의 OCR']
        self.warnings.append('본문은 서식이 없는 검토용 텍스트입니다. 클라우드 버전 이력과 지원하지 않는 개체는 검사 범위에 포함되지 않습니다. 저장 후 원래 문서 프로그램에서도 배치를 확인해 주세요.')
        self.warnings.append('옛 내용이 남을 수 있는 문서 내부 축소판 이미지는 사본에서 제거합니다.')

    @staticmethod
    def package_location(name):
        low = name.lower()
        if 'notes' in low: return f'발표자 노트 · {Path(name).stem}'
        if 'comment' in low or 'memo' in low: return f'메모 · {Path(name).stem}'
        if 'docprops' in low: return '문서 속성'
        if 'header' in low: return '머리말 / 문서 정보'
        if 'footer' in low: return '꼬리말'
        if 'footnote' in low or 'endnote' in low: return '각주·미주'
        if 'slide' in low: return f'슬라이드 · {Path(name).stem}'
        if 'section' in low: return f'본문 · {Path(name).stem}'
        if 'rel' in low: return '연결 정보'
        return f'본문·기록 · {Path(name).stem}'

    def read_excel(self):
        from openpyxl import load_workbook
        entries = read_zip(self.path)
        if 'xl/workbook.xml' not in entries: raise FormatError('올바른 XLSX 파일이 아닙니다.')
        # Reject before a calculation engine can evaluate any external content.
        forbidden = ('externallinks/', 'connections.xml', 'querytables/', 'vbaproject', '/embeddings/', 'threadedcomments/', '_xmlsignatures', 'pivotcache', 'pivottable', '/charts/', '/tables/', '/slicers/', 'metadata.xml')
        if any(any(s in n.lower() for s in forbidden) for n in entries):
            raise FormatError('외부 데이터·차트·그림·피벗·매크로·내장 개체가 있는 XLSX는 아직 지원하지 않습니다. 값만 있는 별도 사본으로 저장해 검사해 주세요.')
        for name, raw in entries.items():
            if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                if any(local_name(e) in {'dataValidations','conditionalFormatting','extLst','autoFilter','controls','oleObjects'} for e in parse_xml(raw).iter()):
                    raise FormatError('조건부 서식·유효성 검사·필터·추가 개체가 있는 XLSX는 아직 지원하지 않습니다. 해당 기능을 제거한 값 전용 사본이 필요합니다.')
            if '/drawings/' in name:
                if not name.lower().endswith('.vml') or any(local_name(e)=='ClientData' and e.get('ObjectType')!='Note' for e in parse_xml(raw).iter()):
                    raise FormatError('그림·도형·양식이 있는 XLSX는 아직 지원하지 않습니다. 값만 있는 사본으로 검사해 주세요.')
            if name.endswith('.rels'):
                if any(e.get('TargetMode') == 'External' for e in parse_xml(raw).iter()):
                    raise FormatError('외부 링크가 있는 XLSX는 재계산하지 않습니다. 연결을 제거한 사본이 필요합니다.')
        wbxml = parse_xml(entries['xl/workbook.xml'])
        if any(local_name(e) == 'definedName' for e in wbxml.iter()):
            raise FormatError('정의된 이름이 있는 XLSX는 아직 지원하지 않습니다. 값 전용 사본으로 다시 검사해 주세요.')
        w = load_workbook(self.path, keep_links=False)
        formulas = []
        allowed = {'SUM','AVERAGE','MIN','MAX','COUNT','COUNTA','IF','IFERROR','AND','OR','NOT','ABS','ROUND','ROUNDUP','ROUNDDOWN','INT','MOD','DATE','YEAR','MONTH','DAY','LEFT','RIGHT','MID','LEN','TRIM','CONCATENATE','CONCAT','TEXT','UPPER','LOWER','SUBSTITUTE','REPLACE','TRUE','FALSE'}
        for sheet_index, sheet in enumerate(w):
            if sheet.max_row*sheet.max_column > 100_000: raise FormatError('시트가 너무 큽니다. 시트당 10만 셀 범위까지 지원합니다.')
            for row in sheet:
                for c in row:
                    if c.data_type != 'f': continue
                    if not isinstance(c.value, str): raise FormatError('배열·공유·동적 수식은 값 전용 사본에서 검사해 주세요.')
                    f = re.sub(r'"(?:[^"]|"")*"', '""', c.value)
                    functions = set(re.findall(r'([A-Za-z_][A-Za-z_0-9.]*)\s*\(', f))
                    if any(x.upper() not in allowed for x in functions) or any(x in f for x in ('[',']','|','://','{','}','#','@')):
                        raise FormatError(f'{sheet.title}!{c.coordinate}: 지원하지 않는 수식 또는 외부 참조입니다. 계산 결과를 값으로 저장한 사본이 필요합니다.')
                    formulas.append((sheet.title, c.coordinate, c.value))
        if formulas:
            recalculated = self.recalculate()
            cached = load_workbook(recalculated, data_only=True, keep_links=False)
            for sh, cell, formula in formulas:
                if sh not in cached: raise FormatError('재계산 후 시트가 일치하지 않습니다.')
                result = cached[sh][cell]
                if result.data_type == 'e' or result.value is None and formula.strip() != '=""':
                    raise FormatError(f'{sh}!{cell}: 계산 결과가 오류이거나 빈 값인지 확인할 수 없습니다. 사본 생성을 중단했습니다.')
                c = w[sh][cell]; c.value = '' if result.value is None else result.value
                if isinstance(c.value, str): c.data_type = 's'
            self.calculation = {'count': len(formulas), 'at': time.strftime('%Y-%m-%d %H:%M:%S'), 'method': 'LibreOffice 강제 재계산 후 값 고정'}
        self.workbook = w
        for sheet_index, sheet in enumerate(w):
            hidden = ' · 숨김 시트' if sheet.sheet_state != 'visible' else ''
            for row in sheet:
                for c in row:
                    loc = f'{sheet.title}!{c.coordinate}{hidden}'
                    if sheet.row_dimensions[c.row].hidden or sheet.column_dimensions[c.column_letter].hidden: loc += ' · 숨김 행/열'
                    if c.value is not None:
                        value = str(c.value)
                        # Handle phone-like numeric masks, not arbitrary Excel format rendering.
                        fmt = c.number_format
                        numeric = isinstance(c.value, (int,float)) and not isinstance(c.value, bool)
                        if numeric and fmt not in {'General', '0', '0.00', '#,##0', '#,##0.00'}:
                            if re.fullmatch(r'0{2,4}[- ]0{3,4}[- ]0{4}', fmt) and float(c.value).is_integer():
                                digits = str(int(c.value)).zfill(fmt.count('0')); it = iter(digits)
                                value = ''.join(next(it) if ch == '0' else ch for ch in fmt)
                            else:
                                self.warnings.append('일부 숫자·날짜 표시 형식은 원시 값으로 검사합니다. 표시 형식의 문자도 별도 검사합니다.')
                        self.add(value, loc, {'kind':'cell','sheet':sheet_index,'cell':c.coordinate})
                        if self.parts and self.parts[-1]['target']=={'kind':'cell','sheet':sheet_index,'cell':c.coordinate}:
                            self.parts[-1].update(sheet=sheet.title,cell=c.coordinate,row=c.row,column=c.column,hidden=sheet.sheet_state!='visible' or bool(sheet.row_dimensions[c.row].hidden) or bool(sheet.column_dimensions[c.column_letter].hidden))
                        if numeric and fmt not in {'General','0','0.00','#,##0','#,##0.00'}:
                            self.add(fmt, loc+' · 표시 형식', {'kind':'format','sheet':sheet_index,'cell':c.coordinate})
                    if c.comment:
                        self.add(c.comment.text, loc+' · 메모', {'kind':'comment','sheet':sheet_index,'cell':c.coordinate})
                        self.add(c.comment.author or '', loc+' · 메모 작성자', {'kind':'comment_author','sheet':sheet_index,'cell':c.coordinate})
            for field in ('oddHeader','oddFooter','evenHeader','evenFooter','firstHeader','firstFooter'):
                for side in ('left','center','right'):
                    text = getattr(getattr(sheet,field),side).text
                    if text: self.add(text,f'{sheet.title} · 인쇄 머리말/꼬리말',{'kind':'header','sheet':sheet_index,'field':field,'side':side})
        for sheet_index,sheet in enumerate(w): self.add(sheet.title,'시트 이름',{'kind':'sheet_name','sheet':sheet_index})
        for key in ('creator','lastModifiedBy','title','subject','description','keywords','category','identifier'):
            value = getattr(w.properties,key,None)
            if value: self.add(value,'문서 속성 · '+key,{'kind':'property','key':key})
        if w.custom_doc_props: self.blockers.append('사용자 지정 문서 속성이 있어 저장을 지원하지 않습니다.')
        self.coverage = ['전체 시트의 셀·메모·작성자', '숨김 시트·행·열', '기본 문서 속성·인쇄 머리말', '지원 수식 강제 재계산·값 고정']
        self.warnings.append('값 전용 배포 사본입니다. 원본 수식은 보존되며, 사본의 값을 바꿔도 다시 계산하지 않습니다. 차트·그림·피벗·외부 연결·정의된 이름은 지원하지 않습니다.')

    def recalculate(self):
        exe = TOOLS/'libreoffice'/'program'/'soffice.com'
        if not exe.is_file(): raise FormatError('값 확인용 LibreOffice가 준비되지 않았습니다. 설치 안내를 확인해 주세요.')
        profile = self.work/'calc-profile'; (profile/'user').mkdir(parents=True, exist_ok=True)
        (profile/'user'/'registrymodifications.xcu').write_text('''<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item><item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item><item oor:path="/org.openoffice.Office.Calc/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>0</value></prop></item></oor:items>''', encoding='utf-8')
        out = self.work/('calculated-'+str(time.time_ns())); out.mkdir()
        r = subprocess.run([str(exe), '-env:UserInstallation='+profile.as_uri(), '--headless', '--norestore','--nodefault','--convert-to','xlsx:Calc MS Excel 2007 XML','--outdir',str(out),str(self.path)], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=60, creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
        target = out/self.path.name
        if r.returncode or not target.exists(): raise FormatError('계산 결과를 확인하지 못했습니다. Excel에서 값으로 저장한 사본으로 다시 검사해 주세요.')
        return target

    def image(self, raw, location, target, *, use_ocr=True):
        if len(self.assets) >= 40: raise FormatError('이미지·PDF 페이지는 파일당 40개까지 지원합니다.')
        im = Image.open(io.BytesIO(raw))
        if getattr(im, 'n_frames', 1) != 1: raise FormatError('여러 프레임을 가진 이미지는 지원하지 않습니다.')
        if im.width*im.height > MAX_PIXELS: raise FormatError('이미지가 너무 큽니다. 2천만 화소 이하로 나누어 주세요.')
        im = ImageOps.exif_transpose(im).convert('RGB')
        aid = f'a{len(self.assets)}'; path = self.work/f'{aid}.png'; im.save(path)
        asset = {'id':aid, 'location':location, 'width':im.width, 'height':im.height, 'target':target, 'file':path.name}
        self.assets.append(asset)
        if use_ocr:
            try:
                from backend.image_ocr import read_image
                for line in read_image(im,'tessdata'):
                    self.add(line['text'],location+' · OCR',{'kind':'raster','asset':aid},boxes=line['boxes'],asset=aid)
                    self.parts[-1]['ocrConfidences']=line['confidences']
            except Exception:
                self.warnings.append(location+': OCR를 읽지 못했습니다. 이미지에서 가릴 영역을 직접 지정해 주세요.')
        self.warnings.append('이미지 OCR은 글자를 놓치거나 잘못 읽을 수 있습니다. 각 이미지를 열어 영역을 확인하고 필요하면 드래그해 직접 가려 주세요.')
        return asset

    def pdf_lines(self, page, location, target, asset=None, scale=(1,1)):
        for block in page.get_text('rawdict')['blocks']:
            for line in block.get('lines', []):
                chars = [c for span in line['spans'] for c in span['chars']]
                if not chars: continue
                text = ''.join(c['c'] for c in chars)
                boxes = [[c['bbox'][0]*scale[0],c['bbox'][1]*scale[1],c['bbox'][2]*scale[0],c['bbox'][3]*scale[1]] for c in chars]
                self.add(text, location, target.copy(), boxes=boxes, asset=asset)

    def read_pdf(self):
        import pymupdf as fitz
        self.pdf = fitz.open(self.path)
        if self.pdf.needs_pass: raise FormatError('암호화된 PDF는 먼저 암호를 해제한 사본이 필요합니다.')
        if len(self.pdf) > 40: raise FormatError('PDF는 파일당 40쪽까지 검사할 수 있습니다.')
        if self.pdf.xref_length()>30000:raise FormatError('PDF 내부 개체가 너무 많습니다. 문서를 나누어 검사해 주세요.')
        for xref in range(1,self.pdf.xref_length()):
            obj=self.pdf.xref_object(xref)
            if re.search(r'/(?:JavaScript|JS|AA|OpenAction|RichMedia|OCProperties|OCG|OCMD)\b',obj):
                self.blockers.append('실행 동작·멀티미디어·숨김 레이어가 있는 PDF는 검사만 지원합니다. 평면화한 사본이 필요합니다.')
                break
        if self.pdf.embfile_count(): self.blockers.append('첨부 파일이 있는 PDF는 검사만 지원합니다. 첨부를 제거한 사본이 필요합니다.')
        for p, page in enumerate(self.pdf):
            if page.rotation: self.blockers.append('회전된 PDF 페이지는 현재 검사만 지원합니다. 페이지 회전을 정규화한 사본이 필요합니다.')
            if page.first_widget: self.blockers.append('서명·양식이 있는 PDF는 검사만 지원합니다. 평면화한 사본이 필요합니다.')
            if page.first_annot: self.blockers.append('주석이 있는 PDF는 검사만 지원합니다. 주석 내용을 별도로 검토하고 제거한 사본이 필요합니다.')
            if page.get_links(): self.blockers.append('링크가 있는 PDF는 검사만 지원합니다. 링크를 제거한 사본이 필요합니다.')
            pix = page.get_pixmap(matrix=fitz.Matrix(1.5,1.5), alpha=False)
            asset = self.image(pix.tobytes('png'), f'{p+1}쪽', {'kind':'pdfpage','page':p,'pageWidth':page.rect.width,'pageHeight':page.rect.height}, use_ocr=not self.verify)
            self.pdf_lines(page, f'{p+1}쪽 · PDF 글자층', {'kind':'pdftext','page':p}, asset=asset['id'], scale=(asset['width']/page.rect.width,asset['height']/page.rect.height))
        self.coverage = ['PDF 글자층', '페이지 전체 이미지 OCR', '선택 위치의 글자·이미지 픽셀 삭제', '사본의 기본 문서 속성·목차 제거 (첨부·양식·주석 등은 저장 제한)']
        self.warnings.append('PDF는 가린 페이지를 새 PDF에 복사해 저장합니다. 목차·문서 속성·접근성 태그·서명 등 문서 부가 기능은 보존하지 않습니다. 이미지와 글자층 후보가 같은 위치에 겹칠 수 있습니다.')
        self.warnings.append('가림 경계와 겹친 주변 문자·도형도 삭제될 수 있습니다. 저장할 실제 사본에서 범위와 배치를 확인해 주세요.')

    def read_hwp(self):
        bundled = TOOLS/'jre'/'bin'/'java.exe'
        choices = sorted(Path('C:/Program Files/Java').glob('jdk*/bin/java.exe'),reverse=True)
        java = str(bundled) if bundled.is_file() else str(choices[0]) if choices else shutil.which('java')
        jar = TOOLS/'hwplib-1.1.11.jar'; out = self.work/'hwp.txt'
        if not java or not jar.exists(): raise FormatError('HWP 읽기 도구(Java·hwplib)가 준비되지 않았습니다. HWPX로 저장한 사본으로 검사해 주세요.')
        compiled = ROOT/'backend'/'HwpText.class'
        # JRE만 있는 설치판은 소스 실행(JDK 필요)을 못 하므로 미리 컴파일한 클래스를 우선 사용한다.
        entry = ['--class-path',os.pathsep.join([str(jar),str(compiled.parent)]),'HwpText'] if compiled.is_file() else ['--class-path',str(jar),str(ROOT/'backend'/'HwpText.java')]
        subprocess.run([java,'-Xmx256m',*entry,str(self.path),str(out)], timeout=45, check=True, stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
        for i,line in enumerate(out.read_text(encoding='utf-8').splitlines()): self.add(line,f'추출 본문 {i+1}행',{'kind':'readonly'})
        self.blockers = ['HWP는 추출 본문 검사만 지원합니다. 메모·변경 기록·이미지 전체 검토와 수정 저장은 보장하지 않습니다. 한글에서 HWPX 사본으로 저장해 다시 검사해 주세요.']
        self.coverage = ['hwplib가 읽은 본문·표 텍스트 (추출 순서)']
        if not self.parts:
            self.blockers.insert(0,'이 HWP에서 읽을 수 있는 본문을 찾지 못했습니다. 개인정보가 없다는 뜻이 아닙니다. 메모만 있는 문서 등은 HWPX로 변환해 검사해 주세요.')

    def edit_part(self, part, new):
        t=part['target']; kind=t['kind']
        if kind=='line': self.lines[t['index']]=new
        elif kind=='csv': self.rows[t['row']][t['col']]=new
        elif kind in {'xml','attr'}:
            nodes=list(self.trees[t['name']].iter())
            if kind=='attr': nodes[t['node']].set(t['attr'],new)
            else:
                # Keep unaffected run text and styling, distributing replacements by original positions.
                old=part['text']; import difflib
                edits=[(a,b,new[c:d]) for op,a,b,c,d in difflib.SequenceMatcher(a=old,b=new,autojunk=False).get_opcodes() if op!='equal']
                offsets=[]; pos=0
                for idx in t['nodes']:
                    value=nodes[idx].text or '';offsets.append((idx,pos,pos+len(value)));pos+=len(value)
                for a,b,value in reversed(edits):
                    first=True
                    for idx,start,end in offsets:
                        if a==b:
                            if not(start<=a<=end): continue
                        elif end<=a or start>=b: continue
                        node=nodes[idx];text=node.text or ''
                        lo=max(a-start,0);hi=min(b-start,len(text))
                        node.text=text[:lo]+(value if first else '')+text[hi:]
                        node.set('{http://www.w3.org/XML/1998/namespace}space','preserve');first=False
                        if a==b: break
                if ''.join(nodes[i].text or '' for i in t['nodes'])!=new: raise FormatError('문서의 글자 위치가 일치하지 않아 저장하지 않았습니다.')
        elif kind=='property': setattr(self.workbook.properties,t['key'],new)
        elif kind=='header': getattr(getattr(self.workbook.worksheets[t['sheet']],t['field']),t['side']).text=new
        elif kind=='sheet_name':
            if not new or len(new)>31 or any(ch in new for ch in '[]:*?/\\') or any(s.title.casefold()==new.casefold() for i,s in enumerate(self.workbook) if i!=t['sheet']): raise FormatError('시트 이름은 중복되지 않는 1~31자여야 하며 []:*?/\\ 문자를 사용할 수 없습니다.')
            self.workbook.worksheets[t['sheet']].title=new
        elif kind in {'cell','format','comment','comment_author'}:
            c=self.workbook.worksheets[t['sheet']][t['cell']]
            if kind=='cell': c.value=new;c.data_type='s'
            elif kind=='format': c.number_format=new
            elif kind=='comment': c.comment.text=new
            else: c.comment.author=new
        else: raise FormatError('이 영역은 글자 직접 수정을 지원하지 않습니다.')

    def export(self, output, edits, regions):
        if self.blockers: raise FormatError(self.blockers[0])
        for pid,new in edits.items():
            p=next((p for p in self.parts if p['id']==pid),None)
            if p is None: raise FormatError('수정 위치를 찾을 수 없습니다.')
            self.edit_part(p,new)
        asset_by_id={a['id']:a for a in self.assets}
        boxes_by_asset={}
        for r in regions:
            a=asset_by_id.get(r['asset'])
            if a is None: raise FormatError('이미지 위치를 찾을 수 없습니다.')
            b=r['box']
            if len(b)!=4 or not all(isinstance(x,(int,float)) and math.isfinite(x) for x in b) or not(0<=b[0]<b[2]<=a['width'] and 0<=b[1]<b[3]<=a['height']):
                raise FormatError('가림 영역이 이미지 범위를 벗어났습니다.')
            boxes_by_asset.setdefault(a['id'],[]).append(b)
        for a in self.assets:
            if a['target']['kind']=='pdfpage': continue
            im=Image.open(self.work/a['file']).convert('RGB'); draw=ImageDraw.Draw(im)
            for b in boxes_by_asset.get(a['id'],[]):
                box=(math.floor(b[0]),math.floor(b[1]),min(im.width,math.ceil(b[2])),min(im.height,math.ceil(b[3])))
                draw.rectangle((box[0],box[1],box[2]-1,box[3]-1),fill='black')
                if im.crop(box).getbbox() is not None: raise FormatError('이미지 가림 픽셀 검증에 실패했습니다.')
            raw=io.BytesIO()
            target=a['target']
            im.save(raw,format='PNG')
            if target['kind']=='image': Path(output).write_bytes(raw.getvalue())
            else:
                old=target['name'];new=old if old.lower().endswith('.png') else old+'.privacy.png'
                if new!=old:
                    if new in self.entries:raise FormatError('내장 이미지 이름이 충돌해 저장하지 않았습니다.')
                    del self.entries[old]
                    for tree in self.trees.values():
                        for e in tree.iter():
                            for attr,value in list(e.attrib.items()):
                                if E.QName(attr).localname.lower() in {'target','href','partname'} and (value.replace('\\','/').endswith('/'+Path(old).name) or value==Path(old).name):
                                    e.set(attr,value[:-len(Path(old).name)]+Path(new).name)
                                    if e.get('media-type'):e.set('media-type','image/png')
                                    if e.get('ContentType'):e.set('ContentType','image/png')
                    content_types=self.trees.get('[Content_Types].xml')
                    if content_types is not None and not any(e.get('Extension')=='png' for e in content_types):
                        E.SubElement(content_types,'{http://schemas.openxmlformats.org/package/2006/content-types}Default',Extension='png',ContentType='image/png')
                self.entries[new]=raw.getvalue()
        if self.kind in {'.docx','.pptx','.hwpx'}:
            # Remove derived Office thumbnails and their relationships/content-type declarations.
            for name in list(self.entries):
                if name.lower().startswith('docprops/thumbnail'): del self.entries[name]
            for tree in self.trees.values():
                for e in list(tree.iter()):
                    if local_name(e)=='Relationship' and e.get('Type','').endswith('/metadata/thumbnail') or local_name(e)=='Override' and e.get('PartName','').lower().startswith('/docprops/thumbnail'):
                        if e.getparent() is not None:e.getparent().remove(e)
            for name,tree in self.trees.items(): self.entries[name]=E.tostring(tree,xml_declaration=True,encoding='UTF-8')
            if self.kind=='.hwpx':
                for name in list(self.entries):
                    if name.startswith('Preview/'): del self.entries[name]
                # Remove package manifest entries pointing at derived previews.
                for name,tree in self.trees.items():
                    removed=False
                    for e in list(tree.iter()):
                        if any('Preview/' in v for v in e.attrib.values()) and e.getparent() is not None: e.getparent().remove(e);removed=True
                    if removed: self.entries[name]=E.tostring(tree,xml_declaration=True,encoding='UTF-8')
            write_zip(output,self.entries)
        elif self.kind=='.xlsx':
            self.workbook.save(output)
            from openpyxl import load_workbook
            check=load_workbook(output)
            if any(c.data_type=='f' for s in check for row in s for c in row): raise FormatError('사본에 수식이 남아 저장하지 않았습니다.')
        elif self.kind=='.pdf':
            import pymupdf as fitz
            for aid,boxes in boxes_by_asset.items():
                a=asset_by_id[aid];t=a['target'];page=self.pdf[t['page']]
                for b in boxes:
                    rect=fitz.Rect(b[0]*t['pageWidth']/a['width'],b[1]*t['pageHeight']/a['height'],b[2]*t['pageWidth']/a['width'],b[3]*t['pageHeight']/a['height'])
                    page.add_redact_annot(rect,fill=(0,0,0))
                page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_PIXELS,graphics=fitz.PDF_REDACT_LINE_ART_REMOVE_IF_TOUCHED,text=fitz.PDF_REDACT_TEXT_REMOVE)
            clean=fitz.open()
            clean.insert_pdf(self.pdf,links=False,annots=False,widgets=False)
            for xref in range(1,clean.xref_length()):
                for key in ('Metadata','PieceInfo'):
                    if clean.xref_get_key(xref,key)[0]!='null':clean.xref_set_key(xref,key,'null')
            clean.set_metadata({});clean.save(output,garbage=4,deflate=True);clean.close()
        elif self.kind in {'.csv','.tsv'}:
            with open(output,'w',encoding='utf-8-sig',newline='') as f:
                csv.writer(f,delimiter=self.delimiter).writerows([[("'"+v if v.lstrip().startswith(('=','+','-','@')) else v) for v in row] for row in self.rows])
        elif self.kind=='.txt': Path(output).write_bytes(''.join(self.lines).encode('utf-8'))
        elif self.kind=='.hwp': raise FormatError('HWP 수정 저장은 지원하지 않습니다.')

